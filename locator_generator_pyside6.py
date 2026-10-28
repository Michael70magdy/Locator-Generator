"""Locator Generator desktop app.

Install: pip install PySide6 beautifulsoup4 selenium openpyxl lxml
Run:     python locator_generator_pyside6.py

The locator-generation engine below is preserved from the original Tkinter project.
"""
import csv
import ast
import re
import sys
from pathlib import Path
from urllib.parse import urlparse

from PySide6.QtCore import Qt, QObject, QThread, Signal, QPoint, QTimer, QUrl
from PySide6.QtGui import QFont, QKeySequence, QShortcut, QPainter, QPolygon, QPalette, QAction, QCursor
from PySide6.QtWidgets import (
    QApplication, QMainWindow, QWidget, QFrame, QLabel, QPushButton,
    QVBoxLayout, QHBoxLayout, QGridLayout, QLineEdit, QPlainTextEdit,
    QTableWidget, QTableWidgetItem, QAbstractItemView, QHeaderView,
    QSplitter, QStackedWidget, QComboBox, QCheckBox, QToolButton, QMenu,
    QProgressBar, QFileDialog, QMessageBox, QGroupBox, QScrollArea,
    QSizePolicy, QTextBrowser, QTreeWidget, QTreeWidgetItem, QTabWidget
)

try:
    from PySide6.QtWebEngineWidgets import QWebEngineView
    HAS_WEBENGINE = True
except ImportError:
    QWebEngineView = None
    HAS_WEBENGINE = False

try:
    from bs4 import BeautifulSoup
except ImportError:
    BeautifulSoup = None
try:
    from selenium.webdriver.common.by import By as SelBy
    from selenium.webdriver.chrome.options import Options as ChromeOptions
    from selenium.webdriver.chrome.webdriver import WebDriver as ChromeWebDriver
    from selenium.webdriver.edge.options import Options as EdgeOptions
    from selenium.webdriver.edge.webdriver import WebDriver as EdgeWebDriver
    HAS_SELENIUM = True
except ImportError:
    HAS_SELENIUM = False
try:
    import openpyxl
    from openpyxl.utils import get_column_letter
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False

TITLE = "Locator Generator"
FRAMEWORKS = ["Java Selenium", "Python Selenium", "Playwright", "Cypress"]
LOCATOR_STRATEGIES = [
    'data-testid', 'data-test', 'data-qa', 'id', 'ID shorthand', 'name',
    'aria-label', 'placeholder', 'title', 'role', 'type',
    'Tag + all classes', 'Single class', 'Class token', 'Tag + type',
    'Exact normalized text', 'Contains text', 'Link text', 'Class + Text combination',
    'Parent ID + Target Tag', 'Parent ID + Target Class',
    'Parent data-testid + Target Tag', 'Parent data-test + Target Tag',
    'Parent data-qa + Target Tag', 'Tag name', 'Relative DOM path',
]
EXAMPLE_HTML = '<div class="form-group">\n  <label for="username">Username</label>\n  <input id="username" name="userName" class="form-control" placeholder="Enter username" type="text">\n</div>'

def xp_literal(s):
    if "'" not in s: return "'" + s + "'"
    if '"' not in s: return '"' + s + '"'
    return "concat(" + ', "\'", '.join("'" + p + "'" for p in s.split("'")) + ")"

def esc(s):
    return s.replace("\\\\", "\\\\\\\\").replace('"', '\\\\"').replace("\\n", "\\\\n").replace("\\r", "\\\\r")

def cls_list(el):
    v = el.get("class", [])
    return v.split() if isinstance(v, str) else list(v)

def rank(score):
    return "Excellent" if score >= 90 else "Strong" if score >= 75 else "Fair" if score >= 55 else "Weak"


def exclude_locator_candidates(candidates, excluded_types=(), excluded_strategies=()):
    types = {value.upper() for value in excluded_types}
    strategies = set(excluded_strategies)
    return [candidate for candidate in candidates
            if candidate.get('kind', '').upper() not in types
            and candidate.get('strategy') not in strategies]


class PersistentCheckMenu(QMenu):
    """Checkable dropdown that stays open while its items are selected."""
    clear_requested = Signal()

    def mouseReleaseEvent(self, event):
        action = self.actionAt(event.pos())
        if event.button() == Qt.MouseButton.LeftButton and action and action.isEnabled():
            if action.isCheckable():
                action.setChecked(not action.isChecked())
                self.setActiveAction(action)
                event.accept()
                return
            if action.text() == 'Clear selections':
                self.clear_requested.emit()
                event.accept()
                return
        super().mouseReleaseEvent(event)

    def keyPressEvent(self, event):
        if event.key() in (Qt.Key.Key_Space, Qt.Key.Key_Return, Qt.Key.Key_Enter):
            action = self.activeAction()
            if action and action.isEnabled():
                if action.isCheckable():
                    action.setChecked(not action.isChecked())
                    self.setActiveAction(action)
                    event.accept()
                    return
                if action.text() == 'Clear selections':
                    self.clear_requested.emit()
                    event.accept()
                    return
        super().keyPressEvent(event)

def make_xpath_snapshot(soup):
    """Create one reusable lxml tree for XPath match counts on this page."""
    try:
        from lxml import html
        return html.fromstring(str(soup))
    except Exception:
        return None


def generate_candidates(soup, el, xpath_root=None):
    tag = el.name.lower()
    candidates, seen = [], set()
    text = " ".join(el.stripped_strings)
    attrs = {}
    for k, v in el.attrs.items():
        attrs[k] = " ".join(v) if isinstance(v, list) else str(v)
    
    def add(kind, strategy, expr, score, note):
        key = (kind, expr)
        if not expr or key in seen: return
        seen.add(key)
        count = None
        if kind == "CSS":
            try: count = len(soup.select(expr))
            except Exception: pass
        elif kind == "XPath" and xpath_root is not None:
            try:
                count = len(xpath_root.xpath(expr))
            except Exception: pass
        # A selector that does not resolve to the target element is not useful
        # to the user. Keep unknown counts (for optional parser support), but
        # never present a confirmed zero-match selector as a candidate.
        if count == 0:
            return
        candidates.append(dict(kind=kind, strategy=strategy, expr=expr, score=score,
                              rank=rank(score), note=note, count=count))

    for key, score, note in [
        ("data-testid", 99, "Dedicated test attribute; highly stable."),
        ("data-test", 98, "Dedicated test attribute; highly stable."),
        ("data-qa", 97, "Dedicated QA attribute; highly stable."),
        ("id", 94, "Excellent if ID is stable and unique."),
        ("name", 91, "Semantic form attribute; stable."),
        ("aria-label", 88, "Accessible label; stable compared to classes."),
        ("placeholder", 81, "User-facing copy; may change."),
        ("title", 76, "Tooltips; may change."),
        ("role", 70, "Role attribute; broad without context."),
        ("type", 58, "Control type; shared across elements.")
    ]:
        if attrs.get(key):
            val = attrs[key]
            add("CSS", key, f'{tag}[{key}="{val.replace(chr(34), chr(92)+chr(34))}"]', score, note)
            add("XPath", key, f"//{tag}[@{key}={xp_literal(val)}]", score, note)
            if key == "id":
                if re.match(r"^[A-Za-z_][\\w-]*$", val):
                    add("CSS", "ID shorthand", "#" + val, 96, "Concise ID selector.")
                add("XPath", "ID shorthand", f"//*[@id={xp_literal(val)}]", 95, "Concise ID-based XPath.")

    classes = cls_list(el)
    safe = []
    for c in classes:
        if not re.match(r"^-?[_a-zA-Z][_a-zA-Z0-9-]*$", c): continue
        safe.append(c)
    
    if safe:
        selector = tag + "".join("." + c for c in safe)
        volatile = any(re.search(r"ng-star-inserted|ng-option-marked|active|selected|css-|sc-", c, re.IGNORECASE) for c in safe)
        add("CSS", "Tag + all classes", selector, 54 if volatile else 73,
            "Volatile framework classes detected." if volatile else "Class-based selector.")
        for c in safe[:4]:
            add("CSS", "Single class", "." + c, 45, "Single class token.")
            add("XPath", "Class token", f"//{tag}[contains(concat(' ', normalize-space(@class), ' '), {xp_literal(' ' + c + ' ')})]", 52, "Class token XPath match.")

    if attrs.get("type"):
        val = attrs["type"]
        add("CSS", "Tag + type", f'{tag}[type="{val}"]', 58, "Tag and type combination.")
        add("XPath", "Tag + type", f"//{tag}[@type={xp_literal(val)}]", 58, "Tag and type combination.")

    if text:
        compact = re.sub(r"\\s+", " ", text).strip()
        if len(compact) <= 100:
            add("XPath", "Exact normalized text", f"//{tag}[normalize-space(.)={xp_literal(compact)}]", 78, "Exact text match.")
            if len(compact) >= 4:
                add("XPath", "Contains text", f"//{tag}[contains(normalize-space(.), {xp_literal(compact[:60])})]", 60, "Partial text match.")
            if tag == "a":
                add("XPath", "Link text", f"//a[normalize-space(.)={xp_literal(compact)}]", 79, "Anchor link text.")
            if safe and len(compact) <= 60:
                add("XPath", "Class + Text combination", f"//{tag}[contains(@class, {xp_literal(safe[0])}) and normalize-space(.)={xp_literal(compact)}]", 86, "High uniqueness class + text combo.")

    parent = el.parent
    if parent and getattr(parent, "name", None) and parent.name != "[document]":
        p_attrs = {}
        for pk, pv in parent.attrs.items():
            p_attrs[pk] = " ".join(pv) if isinstance(pv, list) else str(pv)
        for p_key in ("id", "data-testid", "data-test", "data-qa", "name", "class"):
            if p_val := p_attrs.get(p_key):
                p_val_clean = p_val.strip()
                if p_key == "id":
                    add("CSS", "Parent ID + Target Tag", f"#{p_val_clean} > {tag}", 84, "Scoped under parent ID.")
                    add("XPath", "Parent ID + Target Tag", f"//*[@id={xp_literal(p_val_clean)}]//{tag}", 83, "Scoped under parent ID.")
                    if safe:
                        add("CSS", "Parent ID + Target Class", f"#{p_val_clean} .{safe[0]}", 87, "Parent ID + class combo.")
                    break
                elif p_key in ("data-testid", "data-test", "data-qa"):
                    add("XPath", f"Parent {p_key} + Target Tag", f"//*[@{p_key}={xp_literal(p_val_clean)}]//{tag}", 82, f"Scoped via parent {p_key}.")
                    break

    add("CSS", "Tag name", tag, 30, "Generic tag selector.")
    add("XPath", "Tag name", f"//{tag}", 30, "Generic tag XPath selector.")

    path, node = [], el
    while node and getattr(node, "name", None) and node.name != "[document]" and len(path) < 6:
        parent_node = node.parent
        siblings = parent_node.find_all(node.name, recursive=False) if parent_node and getattr(parent_node, "find_all", None) else [node]
        seg = node.name
        if len(siblings) > 1: seg += f"[{siblings.index(node)+1}]"
        path.insert(0, seg)
        node = parent_node
    if path:
        add("XPath", "Relative DOM path", "/" + "/".join(path), 35, "Brittle structural path.")

    candidates.sort(key=lambda x: (-x["score"], x["kind"], x["strategy"]))
    return candidates



def format_framework_code(candidate, framework):
    """Original four framework bindings, unchanged."""
    expr, kind = candidate['expr'], candidate['kind']
    if framework == 'Java Selenium':
        return f'By.xpath("{esc(expr)}")' if kind == 'XPath' else f'By.cssSelector("{esc(expr)}")'
    if framework == 'Python Selenium':
        by = 'By.XPATH' if kind == 'XPath' else 'By.CSS_SELECTOR'
        return f'driver.find_element({by}, "{esc(expr)}")'
    if framework == 'Playwright':
        return f'page.locator("xpath={esc(expr)}")' if kind == 'XPath' else f'page.locator("{esc(expr)}")'
    if framework == 'Cypress':
        return f'cy.xpath("{esc(expr)}")' if kind == 'XPath' else f'cy.get("{esc(expr)}")'
    return expr


def generate_descriptive_name(expr):
    """Preserved variable-naming heuristics from the Tkinter app."""
    clean_expr = expr.lower()
    base_name = 'element'
    if 'id=' in clean_expr or '#' in clean_expr:
        match = re.search(r'(?:id=["\']?|#)([a-zA-Z0-9\-_]+)', clean_expr)
        if match: base_name = match.group(1)
    elif 'name=' in clean_expr:
        match = re.search(r'name=["\']?([a-zA-Z0-9\-_]+)', clean_expr)
        if match: base_name = match.group(1)
    elif 'data-testid' in clean_expr or 'data-qa' in clean_expr:
        match = re.search(r'data-(?:testid|qa)=["\']?([a-zA-Z0-9\-_]+)', clean_expr)
        if match: base_name = match.group(1)
    elif 'placeholder=' in clean_expr:
        match = re.search(r'placeholder=["\']?([a-zA-Z0-9\-_ ]+)', clean_expr)
        if match: base_name = match.group(1)
    elif 'button' in clean_expr:
        base_name = 'button'
    elif 'input' in clean_expr:
        base_name = 'input'
    elif 'a[' in clean_expr or 'link' in clean_expr:
        base_name = 'link'
    words = re.findall(r'[a-zA-Z0-9]+', base_name)
    if not words: return 'uiElement'
    camel_case = words[0].lower() + ''.join(w.capitalize() for w in words[1:])
    if 'button' in clean_expr and not camel_case.lower().endswith('button'):
        camel_case += 'Button'
    elif 'input' in clean_expr and not camel_case.lower().endswith('input'):
        camel_case += 'Input'
    elif 'a[' in clean_expr and not camel_case.lower().endswith('link'):
        camel_case += 'Link'
    return camel_case


def generate_pom_code(elements, class_name, framework):
    """Keep original POM formats; syntax is the same as the Tkinter edition."""
    if framework == 'Java Selenium':
        lines = [
            'import org.openqa.selenium.By;',
            'import org.openqa.selenium.WebDriver;',
            'import org.openqa.selenium.WebElement;',
            'import org.openqa.selenium.support.PageFactory;',
            '', f'public class {class_name} {{',
            '    private WebDriver driver;', ''
        ]
        for el in elements:
            by_stmt = f'By.xpath("{esc(el["expr"])}")' if el['kind'] == 'XPath' else f'By.cssSelector("{esc(el["expr"])}")'
            lines.append(f'    private final By {el["varname"]} = {by_stmt};')
        lines += ['', f'    public {class_name}(WebDriver driver) {{',
                  '        this.driver = driver;', '    }', '}']
    elif framework == 'Python Selenium':
        lines = ['from selenium.webdriver.common.by import By', '',
                 f'class {class_name}:', '    def __init__(self, driver):',
                 '        self.driver = driver', '']
        for el in elements:
            kind = 'By.XPATH' if el['kind'] == 'XPath' else 'By.CSS_SELECTOR'
            lines.append(f'    {el["varname"]} = ({kind}, "{esc(el["expr"])}")')
        lines.append('')
    elif framework == 'Playwright':
        lines = [f'class {class_name}:', '    def __init__(self, page):',
                 '        self.page = page', '']
        for el in elements:
            value = 'xpath=' + el['expr'] if el['kind'] == 'XPath' else el['expr']
            lines.append(f'        self.{el["varname"]} = page.locator("{esc(value)}")')
        lines.append('')
    else:  # Cypress
        lines = [f'class {class_name} {{']
        for el in elements:
            command = f'cy.xpath("{esc(el["expr"])}")' if el['kind'] == 'XPath' else f'cy.get("{esc(el["expr"])}")'
            lines.append(f'    {el["varname"]} = () => {command};')
        lines.append('}')
    return '\n'.join(lines)


def selenium_driver(headless=True):
    if not HAS_SELENIUM:
        raise RuntimeError('Selenium is not installed. Run: pip install selenium')
    return _start_supported_browser(headless=headless)[0]


def _driver_error_summary(error):
    lines = [line.strip() for line in str(error).splitlines() if line.strip()]
    return ' '.join(lines[:3])[:420] or type(error).__name__


def _start_supported_browser(headless):
    """Start Chrome, falling back to Microsoft Edge if Chrome is unavailable."""
    if not HAS_SELENIUM:
        raise RuntimeError('Selenium is missing from this build.')

    errors = []
    for browser_name, driver_class, options_class in (
            ('Google Chrome', ChromeWebDriver, ChromeOptions),
            ('Microsoft Edge', EdgeWebDriver, EdgeOptions)):
        options = options_class()
        options.add_argument('--no-first-run')
        options.add_argument('--no-default-browser-check')
        if headless:
            options.add_argument('--headless=new')
        try:
            return driver_class(options=options), browser_name
        except Exception as exc:
            errors.append(f'{browser_name}: {_driver_error_summary(exc)}')

    detail = '\n'.join(errors)
    raise RuntimeError(
        'Could not start Google Chrome or Microsoft Edge. Install or repair one of these browsers, '
        'then try again. Selenium Manager may need an internet connection to download the matching browser driver.\n\n'
        + detail
    )


def validate_http_url(value):
    """Return a normalized HTTP(S) URL or raise a readable input error."""
    value = (value or '').strip()
    if not value:
        raise ValueError('Enter a target URL.')
    if '://' not in value:
        value = 'https://' + value
    parsed = urlparse(value)
    if parsed.scheme not in ('http', 'https') or not parsed.netloc:
        raise ValueError('Enter a valid HTTP or HTTPS URL.')
    return value


def scan_html_snapshot(markup, url, excluded_types=(), excluded_strategies=()):
    if BeautifulSoup is None:
        raise RuntimeError('BeautifulSoup is not installed.')
    soup = BeautifulSoup(markup or '', 'html.parser')
    xpath_root = make_xpath_snapshot(soup)
    elements = soup.find_all(['input', 'button', 'a', 'select', 'textarea', 'form'])
    others = soup.find_all(lambda tag: tag.has_attr('id') or tag.has_attr('data-testid') or tag.has_attr('data-qa'))
    all_elements = list(dict.fromkeys(elements + others))
    candidates = []
    seen_locators = set()
    for element in all_elements[:150]:
        matches = exclude_locator_candidates(
            generate_candidates(soup, element, xpath_root),
            excluded_types, excluded_strategies,
        )
        if matches:
            top = matches[0].copy()
            locator_key = (top['kind'], top['expr'])
            if locator_key in seen_locators:
                continue
            seen_locators.add(locator_key)
            top['note'] = f'<{element.name}> | ' + top['note']
            candidates.append(top)
    return {'rows': candidates, 'url': url, 'elements': len(all_elements)}


def scan_page_job(url, excluded_types=(), excluded_strategies=()):
    url = validate_http_url(url)
    driver = selenium_driver()
    try:
        driver.get(url)
        return scan_html_snapshot(driver.page_source, url, excluded_types, excluded_strategies)
    finally:
        driver.quit()


def test_live_job(url, candidates):
    url = validate_http_url(url)
    driver = selenium_driver()
    try:
        driver.get(url)
        results = []
        for candidate in candidates:
            by = SelBy.XPATH if candidate['kind'] == 'XPath' else SelBy.CSS_SELECTOR
            try:
                matches = len(driver.find_elements(by, candidate['expr']))
                results.append((candidate['expr'], matches, None))
            except Exception as exc:
                results.append((candidate['expr'], None, str(exc).splitlines()[0][:200]))
        return {'results': results, 'url': url}
    finally:
        driver.quit()


def generate_locators_job(markup, excluded_types=(), excluded_strategies=()):
    """Parse pasted markup and generate candidates outside the GUI thread."""
    soup = BeautifulSoup(markup, 'html.parser')
    elements = soup.find_all(True)
    if not elements:
        raise ValueError('No HTML elements found. Paste a complete element.')
    preferred_keys = ('data-testid', 'data-test', 'data-qa', 'id', 'name', 'aria-label', 'placeholder')
    target = next((element for element in elements
                   if any(element.get(key) for key in preferred_keys)), elements[-1])
    xpath_root = make_xpath_snapshot(soup)
    rows = exclude_locator_candidates(
        generate_candidates(soup, target, xpath_root), excluded_types, excluded_strategies)
    return {'rows': rows, 'tag': target.name}


def suggest_locator_fallbacks(soup, expr, limit=3, xpath_root=None):
    """Suggest ranked selectors for likely related DOM elements; never claim live validation."""
    if soup is None:
        return []
    if xpath_root is None:
        xpath_root = make_xpath_snapshot(soup)
    hints = {'tag': None, 'id': None, 'classes': [], 'attrs': {}, 'text': None}
    xpath_tag = re.search(r'//([A-Za-z][\w:-]*)', expr)
    css_tag = re.match(r'^\s*([A-Za-z][\w:-]*)', expr)
    tag_match = xpath_tag or css_tag
    if tag_match:
        hints['tag'] = tag_match.group(1).lower()

    id_match = re.search(r'#([A-Za-z_][\w-]*)|@id\s*=\s*([\'\"])(.*?)\2', expr)
    if id_match:
        hints['id'] = id_match.group(1) or id_match.group(3)

    hints['classes'] = re.findall(r'\.([A-Za-z_][\w-]*)', expr)
    class_match = re.search(r'@class\s*=\s*([\'\"])(.*?)\1|contains\s*\(\s*@class\s*,\s*([\'\"])(.*?)\3', expr)
    if class_match:
        value = class_match.group(2) or class_match.group(4) or ''
        hints['classes'].extend(value.split())
    hints['classes'] = list(dict.fromkeys(hints['classes']))

    for name, quote, value in re.findall(r'\[([\w:-]+)\s*=\s*([\'\"])(.*?)\2\]', expr):
        hints['attrs'][name.lower()] = value
    for name, quote, value in re.findall(r'@([\w:-]+)\s*=\s*([\'\"])(.*?)\2', expr):
        if name.lower() != 'id' and name.lower() != 'class':
            hints['attrs'][name.lower()] = value
    text_match = re.search(r'(?:normalize-space\s*\(\s*\.\s*\)|text\s*\(\s*\)|\.)\s*=\s*([\'\"])(.*?)\1|contains\s*\([^,]+,\s*([\'\"])(.*?)\3', expr)
    if text_match:
        hints['text'] = (text_match.group(2) or text_match.group(4) or '').strip()

    if not any((hints['tag'], hints['id'], hints['classes'], hints['attrs'], hints['text'])):
        return []

    scored_elements = []
    for element in soup.find_all(True):
        score, reasons = 0, []
        if hints['tag']:
            if element.name.lower() != hints['tag']:
                continue
            score += 3
            reasons.append(f'tag <{hints["tag"]}>')
        if hints['id']:
            actual = str(element.get('id', ''))
            if actual == hints['id']:
                score += 12
                reasons.append('matching id')
            elif hints['id'].lower() in actual.lower() or actual.lower() in hints['id'].lower():
                score += 5
                reasons.append('similar id')
        actual_classes = set(cls_list(element))
        shared_classes = actual_classes.intersection(hints['classes'])
        if shared_classes:
            score += 4 * len(shared_classes)
            reasons.append('matching class ' + ', '.join(sorted(shared_classes)))
        for name, value in hints['attrs'].items():
            actual = element.get(name)
            if actual is not None:
                actual = ' '.join(actual) if isinstance(actual, list) else str(actual)
                if actual == value:
                    score += 8
                    reasons.append(f'matching {name}')
                elif value.lower() in actual.lower() or actual.lower() in value.lower():
                    score += 3
                    reasons.append(f'similar {name}')
        if hints['text']:
            actual_text = ' '.join(element.stripped_strings)
            if actual_text == hints['text']:
                score += 8
                reasons.append('matching text')
            elif hints['text'].lower() in actual_text.lower() or actual_text.lower() in hints['text'].lower():
                score += 3
                reasons.append('similar text')
        if score:
            scored_elements.append((score, element, reasons))

    scored_elements.sort(key=lambda item: (-item[0], item[1].name))
    suggestions, seen = [], {expr}
    for evidence_score, element, reasons in scored_elements[:30]:
        for candidate in generate_candidates(soup, element, xpath_root):
            selector = candidate['expr']
            if selector in seen:
                continue
            seen.add(selector)
            dom_count = candidate.get('count')
            if dom_count != 1:
                continue
            suggestions.append((evidence_score, candidate['score'], selector,
                                f'DOM snapshot: 1 match; evidence: {", ".join(reasons)}; locator rank: {candidate["rank"]}. Unverified live.'))
    suggestions.sort(key=lambda item: (-item[0], -item[1], item[2]))
    return [(selector, note) for _, _, selector, note in suggestions[:limit]]


def audit_job(url, locators):
    url = validate_http_url(url)
    driver = selenium_driver()
    try:
        driver.get(url)
        soup = BeautifulSoup(driver.page_source, 'html.parser') if BeautifulSoup else None
        xpath_root = make_xpath_snapshot(soup) if soup is not None else None
        results, suggestions = [], []
        for expr in locators:
            by = SelBy.XPATH if expr.startswith(('//', '(', './/', '/', 'id(')) else SelBy.CSS_SELECTOR
            try:
                matches = len(driver.find_elements(by, expr))
                status = 'Found' if matches else '0 matches'
                results.append((expr, str(matches), status))
            except Exception as error:
                # Never claim a broken locator when the underlying call could not evaluate it.
                results.append((expr, '—', 'Evaluation error'))
                suggestions.append((expr, 'Cannot evaluate selector', str(error).split('\n')[0][:160]))
                continue
            if matches:
                continue
            fallbacks = suggest_locator_fallbacks(soup, expr, xpath_root=xpath_root)
            if fallbacks:
                suggestions.extend((expr, selector, note) for selector, note in fallbacks)
            else:
                suggestions.append((expr, 'No likely fallback found',
                                    'The page snapshot did not contain a sufficiently similar, unique element.'))
        return {'results': results, 'suggestions': suggestions, 'url': url}
    finally:
        driver.quit()


def extract_audit_locators(source):
    """Extract selectors from plain selectors or common Selenium declarations.

    Accepts forms such as ``By login = By.id("login")``,
    ``login = (By.CSS_SELECTOR, "#login")`` and direct CSS/XPath lines.
    Selenium's ID/name/class/tag strategies are converted to CSS selectors so
    the audit runner can consistently evaluate CSS and XPath.
    """
    if not source:
        return []

    def decode(value):
        try:
            return str(ast.literal_eval(value))
        except (ValueError, SyntaxError):
            return value[1:-1] if len(value) >= 2 else value

    def normalize(strategy, value):
        strategy = strategy.upper().split('.')[-1]
        if strategy in ('XPATH', 'LINK_TEXT', 'PARTIAL_LINK_TEXT'):
            if strategy == 'LINK_TEXT':
                return f"//a[normalize-space(.)={xp_literal(value)}]"
            if strategy == 'PARTIAL_LINK_TEXT':
                return f"//a[contains(normalize-space(.), {xp_literal(value)})]"
            return value
        if strategy in ('ID',):
            return '#' + value if re.fullmatch(r'[A-Za-z_][\w-]*', value) else f'[id="{value}"]'
        if strategy == 'NAME':
            return f'[name="{value}"]'
        if strategy == 'CLASS_NAME':
            return '.' + value if re.fullmatch(r'-?[_a-zA-Z][_a-zA-Z0-9-]*', value) else f'[class~="{value}"]'
        if strategy == 'TAG_NAME':
            return value
        return value

    extracted = []
    for raw_line in source.splitlines():
        line = raw_line.strip().rstrip(';').strip()
        if not line or line == '//' or line.startswith(('// ', '//\t', '# ', '#\t')):
            continue

        # Java: By.variable = By.id("..."); Python calls/tuples: By.ID, "..."
        by_call = re.search(r'By\.(xpath|cssSelector|id|name|className|tagName|linkText|partialLinkText)\s*\(\s*((?:"(?:\\.|[^"\\])*"|\'(?:\\.|[^\'\\])*\'))\s*\)', line, re.I)
        if by_call:
            # Java Selenium uses camelCase names, while normalize() expects
            # Selenium's underscore-separated strategy constants.
            strategy = {
                'cssselector': 'CSS_SELECTOR',
                'classname': 'CLASS_NAME',
                'tagname': 'TAG_NAME',
                'linktext': 'LINK_TEXT',
                'partiallinktext': 'PARTIAL_LINK_TEXT',
            }.get(by_call.group(1).lower(), by_call.group(1))
            extracted.append(normalize(strategy, decode(by_call.group(2))))
            continue

        tuple_call = re.search(r'By\.(XPATH|CSS_SELECTOR|ID|NAME|CLASS_NAME|TAG_NAME|LINK_TEXT|PARTIAL_LINK_TEXT)\s*,\s*((?:"(?:\\.|[^"\\])*"|\'(?:\\.|[^\'\\])*\'))', line, re.I)
        if tuple_call:
            extracted.append(normalize(tuple_call.group(1), decode(tuple_call.group(2))))
            continue

        # Locator APIs commonly wrap the selector in a quoted call argument.
        wrapped = re.search(r'(?:page\.locator|cy\.get|cy\.xpath)\s*\(\s*((?:"(?:\\.|[^"\\])*"|\'(?:\\.|[^\'\\])*\'))', line, re.I)
        if wrapped:
            value = decode(wrapped.group(1))
            extracted.append(value[6:] if value.startswith('xpath=') else value)
            continue

        # Direct selectors may also be assigned to variables. Extract the
        # quoted right-hand-side only when it looks like a CSS/XPath selector.
        assignment = re.search(r'=\s*((?:"(?:\\.|[^"\\])*"|\'(?:\\.|[^\'\\])*\'))\s*$', line)
        candidate = decode(assignment.group(1)) if assignment else line
        if candidate.startswith('xpath='):
            candidate = candidate[6:]
        if candidate.startswith(('//', './/', '/', '(', 'id(')) or candidate.startswith(('#', '.', '[', '*')):
            extracted.append(candidate)
        elif re.fullmatch(r'[A-Za-z][\w-]*', candidate):
            extracted.append(candidate)

    # Preserve order while dropping duplicates and blank values.
    return list(dict.fromkeys(value.strip() for value in extracted if value.strip()))


# Browser automation runs exclusively in its own worker thread. GUI events only
# enqueue commands; they never call into Selenium from the Qt GUI thread.
PICKER_SCRIPT = r"""
if (!window.__locatorPickerInstalled) {
  window.__locatorPickerInstalled = true;
  window.__selectedElementHtml = null;
  const overlay = document.createElement('div');
  overlay.style.cssText = 'position:fixed;z-index:2147483646;border:2px solid #2563eb;background:rgba(37,99,235,.12);pointer-events:none;box-sizing:border-box';
  document.documentElement.appendChild(overlay);
  window.__locatorPickerOverlay = overlay;
  document.addEventListener('mouseover', function(e) {
    if (!window.__pickerEnabled) { overlay.style.display = 'none'; return; }
    overlay.style.display = 'block';
    const rect = e.target.getBoundingClientRect();
    overlay.style.top = rect.top+'px'; overlay.style.left = rect.left+'px';
    overlay.style.width = rect.width+'px'; overlay.style.height = rect.height+'px';
  }, true);
  document.addEventListener('click', function(e) {
    if (!window.__pickerEnabled) return;
    e.preventDefault(); e.stopPropagation();
    const target = e.target.closest('button, a, input, select, textarea, label, [role="button"], [role="link"], [data-testid], [data-test], [data-qa]') || e.target;
    window.__selectedElementHtml = target.outerHTML;
  }, true);
}
window.__pickerEnabled = PICKER_ENABLED;
"""


class Worker(QObject):
    result = Signal(object)
    error = Signal(str)
    finished = Signal()

    def __init__(self, func):
        super().__init__()
        self.func = func

    def run(self):
        try:
            self.result.emit(self.func())
        except Exception as exc:
            self.error.emit(str(exc))
        finally:
            self.finished.emit()


def label(text, secondary=False):
    obj = QLabel(text)
    if secondary:
        obj.setObjectName('muted')
    return obj


def panel(title, content=None):
    frame = QFrame()
    frame.setObjectName('panel')
    layout = QVBoxLayout(frame)
    layout.setContentsMargins(14, 12, 14, 12)
    layout.setSpacing(10)
    if title:
        name = QLabel(title)
        name.setObjectName('sectionTitle')
        layout.addWidget(name)
    if content:
        layout.addWidget(content, 1)
    return frame, layout


def table(headers, widths=None):
    view = QTableWidget(0, len(headers))
    view.setHorizontalHeaderLabels(headers)
    view.setSelectionBehavior(QAbstractItemView.SelectionBehavior.SelectRows)
    view.setSelectionMode(QAbstractItemView.SelectionMode.ExtendedSelection)
    view.setEditTriggers(QAbstractItemView.EditTrigger.NoEditTriggers)
    view.setAlternatingRowColors(True)
    view.verticalHeader().setVisible(False)
    view.setWordWrap(False)
    view.setHorizontalScrollMode(QAbstractItemView.ScrollMode.ScrollPerPixel)
    view.horizontalHeader().setStretchLastSection(True)
    view.horizontalHeader().setDefaultAlignment(Qt.AlignmentFlag.AlignLeft)
    for index, width in enumerate(widths or []):
        view.setColumnWidth(index, width)
    return view


def put_rows(view, rows):
    view.setUpdatesEnabled(False)
    try:
        view.setRowCount(0)
        for row in rows:
            at = view.rowCount()
            view.insertRow(at)
            for col, value in enumerate(row):
                item = QTableWidgetItem(str(value))
                if col == 0:
                    item.setToolTip(str(value))
                view.setItem(at, col, item)
    finally:
        view.setUpdatesEnabled(True)


class ArrowComboBox(QComboBox):
    """Combo box with a clearly drawn down arrow and the standard popup list."""
    def paintEvent(self, event):
        super().paintEvent(event)
        painter = QPainter(self)
        painter.setRenderHint(QPainter.RenderHint.Antialiasing, True)
        painter.setPen(Qt.PenStyle.NoPen)
        painter.setBrush(self.palette().color(QPalette.ColorRole.ButtonText))
        center_x = self.width() - 15
        center_y = self.height() // 2
        arrow = QPolygon([
            QPoint(center_x - 4, center_y - 2),
            QPoint(center_x + 4, center_y - 2),
            QPoint(center_x, center_y + 3),
        ])
        painter.drawPolygon(arrow)


class ArrowMenuButton(QToolButton):
    """Menu button styled with the same single down-arrow as ArrowComboBox."""
    def paintEvent(self, event):
        super().paintEvent(event)
        painter = QPainter(self)
        painter.setRenderHint(QPainter.RenderHint.Antialiasing, True)
        painter.setPen(Qt.PenStyle.NoPen)
        painter.setBrush(self.palette().color(QPalette.ColorRole.ButtonText))
        center_x = self.width() - 15
        center_y = self.height() // 2
        painter.drawPolygon(QPolygon([
            QPoint(center_x - 4, center_y - 2),
            QPoint(center_x + 4, center_y - 2),
            QPoint(center_x, center_y + 3),
        ]))


class MainWindow(QMainWindow):
    job_error = Signal(str)
    job_finished = Signal()

    def __init__(self):
        super().__init__()
        self.setWindowTitle(TITLE)
        screen = QApplication.primaryScreen()
        available = screen.availableGeometry() if screen else None
        initial_width = min(1440, max(420, available.width() - 32)) if available else 1200
        initial_height = min(900, max(360, available.height() - 48)) if available else 760
        self.resize(initial_width, initial_height)
        self.setMinimumSize(min(420, initial_width), min(360, initial_height))
        self.ui_scale = 1.0
        self._responsive_splitters = []
        self._workspace_layouts = []
        self._workspace_subtitles = []
        self._workspace_titles = []
        self.rows = []
        self.pom_elements = []
        self.count_context = 'HTML fragment'
        self.dark = True
        self.busy = False
        self.worker_threads = []  # Strong references until Qt threads finish
        self._job_success = None
        self.captured_bundles = []
        self._pending_capture_ids = []
        self._active_capture_id = None
        self._next_capture_id = 1
        self._next_capture_number = 1
        self.loaded_bundle_id = None
        self._picker_poll_in_flight = False
        self._picker_loaded = False
        self._picker_page_open = False
        self._picker_live_results = {}
        self.picker_poll_timer = None
        self._close_when_idle = False
        self._workspace_drawer_open = True
        self.framework = FRAMEWORKS[0]
        self.nav_buttons = []
        self.job_error.connect(self.on_job_error)
        self.job_finished.connect(self.on_job_finished)
        self._setup_ui()
        self.apply_theme()
        self.switch_workspace(0)
        self.set_status('Ready. Paste HTML, pick an element, or scan a page.')

    def _setup_ui(self):
        base = QWidget()
        base.setObjectName('root')
        self.setCentralWidget(base)
        root = QVBoxLayout(base)
        root.setContentsMargins(0, 0, 0, 0)
        root.setSpacing(0)

        header = QFrame()
        header.setObjectName('header')
        head = QHBoxLayout(header)
        head.setContentsMargins(22, 12, 22, 12)
        self.header_layout = head
        title_box = QVBoxLayout()
        title_box.setSpacing(2)
        title = QLabel('Locator Generator')
        title.setObjectName('appTitle')
        title_box.addWidget(title)
        self.header_subtitle = label('QA automation workspace  /  CSS · XPath · Page Objects · Audits', True)
        title_box.addWidget(self.header_subtitle)
        head.addLayout(title_box, 1)
        self.responsive_nav = ArrowComboBox()
        workspace_names = ['Element Picker', 'Page Objects', 'Health Audit', 'Help']
        self.responsive_nav.addItems(workspace_names)
        self.responsive_nav.setMinimumWidth(150)
        self.responsive_nav.currentIndexChanged.connect(self.switch_workspace)
        self.responsive_nav.hide()
        head.addWidget(self.responsive_nav)
        self.zoom_out_button = QPushButton('A−')
        self.zoom_out_button.setObjectName('zoomButton')
        self.zoom_out_button.setToolTip('Zoom out (Ctrl+-)')
        self.zoom_out_button.clicked.connect(lambda: self.change_zoom(-0.1))
        head.addWidget(self.zoom_out_button)
        self.zoom_label = label('100%', True)
        self.zoom_label.setMinimumWidth(38)
        self.zoom_label.setAlignment(Qt.AlignmentFlag.AlignCenter)
        head.addWidget(self.zoom_label)
        self.zoom_in_button = QPushButton('A+')
        self.zoom_in_button.setObjectName('zoomButton')
        self.zoom_in_button.setToolTip('Zoom in (Ctrl++)')
        self.zoom_in_button.clicked.connect(lambda: self.change_zoom(0.1))
        head.addWidget(self.zoom_in_button)
        self.theme_button = QPushButton('Light theme')
        self.theme_button.setObjectName('secondaryButton')
        self.theme_button.setToolTip('Switch between dark and light theme.')
        self.theme_button.clicked.connect(self.toggle_theme)
        head.addWidget(self.theme_button)
        root.addWidget(header)

        middle = QHBoxLayout()
        middle.setSpacing(0)
        root.addLayout(middle, 1)
        sidebar = QFrame()
        sidebar.setObjectName('sidebar')
        sidebar.setFixedWidth(194)
        self.sidebar = sidebar
        side = QVBoxLayout(sidebar)
        side.setContentsMargins(12, 20, 12, 16)
        side.setSpacing(7)
        side.addWidget(label('WORKSPACES', True))
        for i, title in enumerate(workspace_names):
            b = QPushButton(title)
            b.setObjectName('navButton')
            b.setCheckable(True)
            b.setMinimumHeight(40)
            b.clicked.connect(lambda checked=False, index=i: self.switch_workspace(index))
            self.nav_buttons.append(b)
            side.addWidget(b)
        side.addStretch(1)
        middle.addWidget(sidebar)

        self.workspace_drawer_button = QToolButton()
        self.workspace_drawer_button.setObjectName('drawerButton')
        self.workspace_drawer_button.setFixedWidth(30)
        self.workspace_drawer_button.setText('◀')
        self.workspace_drawer_button.setToolTip('Close the workspace drawer.')
        self.workspace_drawer_button.clicked.connect(self.toggle_workspace_drawer)
        middle.addWidget(self.workspace_drawer_button, 0, Qt.AlignmentFlag.AlignVCenter)

        self.workspaces = QStackedWidget()
        self.workspaces.setObjectName('workspaces')
        middle.addWidget(self.workspaces, 1)
        self.build_exclusion_menu()
        self.build_picker_workspace()
        self.build_pom_workspace()
        self.build_audit_workspace()
        self.build_help_workspace()

        footer = QFrame()
        footer.setObjectName('footer')
        fl = QHBoxLayout(footer)
        fl.setContentsMargins(18, 7, 18, 7)
        self.status = label('', True)
        fl.addWidget(self.status, 1)
        self.progress = QProgressBar()
        self.progress.setFixedWidth(170)
        self.progress.setFixedHeight(9)
        self.progress.setRange(0, 0)
        self.progress.hide()
        fl.addWidget(self.progress)
        root.addWidget(footer)
        self.shortcuts = []
        for index in range(len(workspace_names)):
            shortcut = QShortcut(QKeySequence(f'Ctrl+{index + 1}'), self)
            shortcut.activated.connect(lambda index=index: self.switch_workspace(index))
            self.shortcuts.append(shortcut)
        zoom_in = QShortcut(QKeySequence('Ctrl++'), self)
        zoom_in.activated.connect(lambda: self.change_zoom(0.1))
        zoom_out = QShortcut(QKeySequence('Ctrl+-'), self)
        zoom_out.activated.connect(lambda: self.change_zoom(-0.1))
        zoom_reset = QShortcut(QKeySequence('Ctrl+0'), self)
        zoom_reset.activated.connect(self.reset_zoom)
        self.shortcuts.extend((zoom_in, zoom_out, zoom_reset))

    def register_responsive_splitter(self, splitter, breakpoint=1080, ratio=0.44, vertical_ratio=None):
        splitter.setChildrenCollapsible(False)
        splitter.setProperty('firstPaneRatio', ratio)
        splitter.setProperty('verticalPaneRatio', vertical_ratio if vertical_ratio is not None else ratio)
        self._responsive_splitters.append((splitter, breakpoint))

    def change_zoom(self, delta):
        self.ui_scale = min(1.5, max(0.5, round(self.ui_scale + delta, 2)))
        self.apply_theme()

    def reset_zoom(self):
        self.ui_scale = 1.0
        self.apply_theme()

    def toggle_workspace_drawer(self):
        self._workspace_drawer_open = not self._workspace_drawer_open
        compact = self.width() < 980 * self.ui_scale
        if compact:
            self.responsive_nav.setVisible(self._workspace_drawer_open)
        else:
            self.sidebar.setVisible(self._workspace_drawer_open)
        self.workspace_drawer_button.setText('◀' if self._workspace_drawer_open else '▶')
        self.workspace_drawer_button.setToolTip(
            'Close the workspace drawer.' if self._workspace_drawer_open
            else 'Open the workspace drawer.'
        )

    def resizeEvent(self, event):
        super().resizeEvent(event)
        width = event.size().width()
        height = event.size().height()
        compact = width < 980 * self.ui_scale
        short = height < 660 * self.ui_scale
        if hasattr(self, 'sidebar'):
            self.sidebar.setVisible(not compact and self._workspace_drawer_open)
        if hasattr(self, 'workspace_drawer_button'):
            self.workspace_drawer_button.setText('◀' if self._workspace_drawer_open else '▶')
            self.workspace_drawer_button.setToolTip(
                'Close the workspace drawer.' if self._workspace_drawer_open
                else 'Open the workspace drawer.'
            )
        if hasattr(self, 'responsive_nav'):
            self.responsive_nav.setVisible(compact and self._workspace_drawer_open)
            self.responsive_nav.setMinimumWidth(110 if width < 700 * self.ui_scale else 150)
            if compact and self.responsive_nav.currentIndex() != self.workspaces.currentIndex():
                self.responsive_nav.blockSignals(True)
                self.responsive_nav.setCurrentIndex(self.workspaces.currentIndex())
                self.responsive_nav.blockSignals(False)
        if hasattr(self, 'header_subtitle'):
            self.header_subtitle.setVisible(width >= 1050 * self.ui_scale and not short)
        if hasattr(self, 'zoom_label'):
            self.zoom_label.setVisible(width >= 540 * self.ui_scale)
        if hasattr(self, 'header_layout'):
            margin_x = 10 if width < 700 * self.ui_scale else 16
            margin_y = 5 if short else 9
            self.header_layout.setContentsMargins(margin_x, margin_y, margin_x, margin_y)
        if hasattr(self, 'theme_button'):
            self.theme_button.setText('Theme' if width < 700 * self.ui_scale else ('Light theme' if self.dark else 'Dark theme'))
        for title, subtitle, layout in zip(self._workspace_titles, self._workspace_subtitles, self._workspace_layouts):
            subtitle.setVisible(not short)
            margin = 8 if short else 14
            layout.setContentsMargins(margin, margin, margin, margin)
            layout.setSpacing(8 if short else 11)
        for splitter, breakpoint in getattr(self, '_responsive_splitters', []):
            effective_breakpoint = breakpoint * self.ui_scale
            wanted = Qt.Orientation.Vertical if width < effective_breakpoint else Qt.Orientation.Horizontal
            if splitter.orientation() != wanted:
                splitter.setOrientation(wanted)
                extent = splitter.height() if wanted == Qt.Orientation.Vertical else splitter.width()
                ratio_property = 'verticalPaneRatio' if wanted == Qt.Orientation.Vertical else 'firstPaneRatio'
                ratio = float(splitter.property(ratio_property) or 0.44)
                splitter.setSizes([max(1, round(extent * ratio)), max(1, round(extent * (1 - ratio)))])

    def workspace(self, title, subtitle):
        container = QWidget()
        outer = QVBoxLayout(container)
        outer.setContentsMargins(18, 16, 18, 16)
        outer.setSpacing(12)
        h = QLabel(title)
        h.setObjectName('pageTitle')
        outer.addWidget(h)
        subtitle_label = label(subtitle, True)
        outer.addWidget(subtitle_label)
        self._workspace_titles.append(h)
        self._workspace_subtitles.append(subtitle_label)
        self._workspace_layouts.append(outer)
        self.workspaces.addWidget(container)
        return outer

    def build_generate_workspace(self):
        outer = self.workspace('Generate Locators', 'Choose an input mode, generate ranked selectors, then export or reuse them.')
        primary_split = QSplitter(Qt.Orientation.Horizontal)
        self.register_responsive_splitter(primary_split, 1210, ratio=0.29, vertical_ratio=0.42)
        outer.addWidget(primary_split, 1)

        input_panel, il = panel('INPUT')
        input_panel.setMinimumWidth(210)
        primary_split.addWidget(input_panel)
        self.mode_buttons = []
        modebar = QHBoxLayout()
        modebar.setSpacing(4)
        for i, text in enumerate(['Paste HTML', 'Scan Page']):
            button = QPushButton(text)
            button.setObjectName('modeButton')
            button.setMinimumHeight(38)
            button.setSizePolicy(QSizePolicy.Policy.Expanding, QSizePolicy.Policy.Fixed)
            button.setToolTip([
                'Paste HTML markup to generate selectors for an element.',
                'Read eligible elements from a live page URL.',
            ][i])
            button.setCheckable(True)
            button.clicked.connect(lambda checked=False, index=i: self.switch_mode(index))
            self.mode_buttons.append(button)
            modebar.addWidget(button)
        il.addLayout(modebar)
        self.modes = QStackedWidget()
        il.addWidget(self.modes, 1)

        paste = QWidget()
        pv = QVBoxLayout(paste)
        pv.setContentsMargins(0, 5, 0, 0)
        pv.setSpacing(9)
        pv.addWidget(label('Paste a target HTML element or a small HTML fragment.', True))
        self.html = QPlainTextEdit()
        self.html.setObjectName('codeArea')
        self.html.setPlaceholderText('Paste an HTML element here, for example:\n<button data-testid="login">Log in</button>')
        self.html.setPlainText(EXAMPLE_HTML)
        pv.addWidget(self.html, 1)
        html_actions = QHBoxLayout()
        paste_clipboard = QPushButton('Paste from clipboard')
        paste_clipboard.setToolTip('Insert HTML text from the system clipboard into the editor.')
        paste_clipboard.clicked.connect(lambda: self.html.setPlainText(QApplication.clipboard().text()))
        html_actions.addWidget(paste_clipboard)
        example = QPushButton('Load example')
        example.setToolTip('Replace the editor contents with a sample HTML element.')
        example.clicked.connect(lambda: self.html.setPlainText(EXAMPLE_HTML))
        html_actions.addWidget(example)
        clear = QPushButton('Clear input')
        clear.setToolTip('Clear the HTML editor and current generated selectors.')
        clear.clicked.connect(self.clear_html)
        html_actions.addWidget(clear)
        html_actions.addStretch(1)
        pv.addLayout(html_actions)
        self.modes.addWidget(paste)

        scan = QWidget()
        sl = QVBoxLayout(scan)
        sl.setContentsMargins(0, 5, 0, 0)
        sl.setSpacing(12)
        sl.addWidget(label('Scan buttons, inputs, links, forms and tagged elements from a live page.', True))
        sl.addWidget(label('Target URL'))
        self.scan_url = QLineEdit('https://example.com')
        sl.addWidget(self.scan_url)
        self.scan_url.editingFinished.connect(lambda: self.live_test_url.setText(self.scan_url.text().strip()))
        self.scan_button = QPushButton('Scan whole page')
        self.scan_button.setObjectName('primaryButton')
        self.scan_button.clicked.connect(self.scan_whole_page)
        sl.addWidget(self.scan_button)
        scan_summary, scan_summary_layout = panel('SCAN SNAPSHOT')
        self.scan_summary = label('No page scanned yet.\n\nA scan collects buttons, inputs, links, forms, and elements with test attributes. Results appear in Candidate Locators.', True)
        self.scan_summary.setWordWrap(True)
        scan_summary_layout.addWidget(self.scan_summary)
        scan_summary.setMinimumHeight(110)
        sl.addWidget(scan_summary, 1)
        scan_note = label('Snapshot match counts are estimates from captured HTML. Use Test live to verify selectors in the browser.', True)
        scan_note.setWordWrap(True)
        sl.addWidget(scan_note)
        sl.addStretch(1)
        self.modes.addWidget(scan)

        self.generate_button = QPushButton('Generate Locators')
        self.generate_button.setObjectName('primaryButton')
        self.generate_button.setMinimumHeight(39)
        self.generate_button.clicked.connect(self.generate)
        il.addWidget(self.generate_button)

        right = QSplitter(Qt.Orientation.Vertical)
        right.setChildrenCollapsible(False)
        primary_split.addWidget(right)
        primary_split.setSizes([290, 895])

        results_panel, rl = panel('CANDIDATE LOCATORS')
        self.exclusion_actions = []
        self.exclusion_menu = PersistentCheckMenu(self)
        self.exclusion_menu.setMinimumWidth(290)
        type_heading = QAction('LOCATOR TYPES', self.exclusion_menu)
        type_heading.setEnabled(False)
        type_font = type_heading.font()
        type_font.setBold(True)
        type_heading.setFont(type_font)
        self.exclusion_menu.addAction(type_heading)
        for title, option_type, value in [('CSS selectors', 'type', 'CSS'), ('XPath selectors', 'type', 'XPath')]:
            action = QAction(title, self.exclusion_menu)
            action.setCheckable(True)
            action.setData((option_type, value))
            action.toggled.connect(self.update_exclusion_label)
            self.exclusion_menu.addAction(action)
            self.exclusion_actions.append((action, option_type, value, title))
        self.exclusion_menu.addSeparator()
        strategy_heading = QAction('LOCATOR STRATEGIES', self.exclusion_menu)
        strategy_heading.setEnabled(False)
        strategy_font = strategy_heading.font()
        strategy_font.setBold(True)
        strategy_heading.setFont(strategy_font)
        self.exclusion_menu.addAction(strategy_heading)
        for strategy in LOCATOR_STRATEGIES:
            action = QAction(strategy, self.exclusion_menu)
            action.setCheckable(True)
            action.setData(('strategy', strategy))
            action.toggled.connect(self.update_exclusion_label)
            self.exclusion_menu.addAction(action)
            self.exclusion_actions.append((action, 'strategy', strategy, strategy))
        self.exclusion_menu.addSeparator()
        clear_exclusions_action = QAction('Clear selections', self.exclusion_menu)
        self.exclusion_menu.addAction(clear_exclusions_action)
        self.exclusion_menu.clear_requested.connect(self.clear_exclusions)
        self.exclude_button = ArrowMenuButton()
        self.exclude_button.setObjectName('excludeButton')
        self.exclude_button.setText('Ignore patterns')
        self.exclude_button.setToolTip('Choose locator types or strategies to omit from the next generated or scanned results. The menu stays open while you choose multiple items.')
        self.exclude_button.setPopupMode(QToolButton.ToolButtonPopupMode.InstantPopup)
        self.exclude_button.setMenu(self.exclusion_menu)
        results_tools = QHBoxLayout()
        results_tools.addWidget(self.exclude_button)
        self.excluded_summary = label('Ignored: none', True)
        self.excluded_summary.setWordWrap(True)
        results_tools.addWidget(self.excluded_summary, 1)
        rl.addLayout(results_tools)
        self.results_caption = label('No locators generated yet.', True)
        rl.addWidget(self.results_caption)
        self.tree = table(['Select', 'Selector', 'Type', 'Strategy', 'Score', 'Source matches'], [55, 385, 72, 172, 98, 120])
        self.tree.itemSelectionChanged.connect(self.preview)
        self.tree.itemChanged.connect(self.on_candidate_check_changed)
        self.tree.itemDoubleClicked.connect(lambda *_: self.copy_expr())
        rl.addWidget(self.tree, 1)
        self.select_all_candidates = QCheckBox('Select all generated selectors')
        self.select_all_candidates.toggled.connect(self.set_all_candidates_checked)
        rl.addWidget(self.select_all_candidates)
        live_row = QHBoxLayout()
        live_row.addWidget(label('Live test URL'))
        self.live_test_url = QLineEdit('https://example.com')
        self.live_test_url.setPlaceholderText('https://your-test-environment.example')
        live_row.addWidget(self.live_test_url, 1)
        rl.addLayout(live_row)
        actions = QHBoxLayout()
        for text, callback in [
            ('Copy selectors', self.copy_expr),
            ('Add to Page Objects', self.add_selected_to_pom),
            ('Send to Health Audit', self.send_selected_to_auditor),
            ('Test live', self.test_live_url),
            ('Export', self.export_excel),
        ]:
            b = QPushButton(text)
            b.clicked.connect(callback)
            actions.addWidget(b)
            if text == 'Test live':
                self.test_button = b
                self.test_button.setEnabled(HAS_SELENIUM)
        rl.addLayout(actions)
        right.addWidget(results_panel)

        lower_split = QTabWidget()
        lower_split.setDocumentMode(True)
        self.generate_detail_tabs = lower_split
        notes_panel, nl = panel('SELECTED LOCATOR')
        self.selected_note = QPlainTextEdit()
        self.selected_note.setReadOnly(True)
        self.selected_note.setPlaceholderText('Select a locator to see its stability notes and match scope.')
        nl.addWidget(self.selected_note, 1)
        lower_split.addTab(notes_panel, 'Locator Details')
        code_panel, cl = panel('CODE PREVIEW')
        options = QHBoxLayout()
        options.addWidget(label('Framework'))
        self.framework_combo = ArrowComboBox()
        self.framework_combo.addItems(FRAMEWORKS)
        self.framework_combo.currentTextChanged.connect(self.change_framework)
        options.addWidget(self.framework_combo, 1)
        self.copy_code_button = QPushButton('Copy code')
        self.copy_code_button.clicked.connect(self.copy_framework_code)
        options.addWidget(self.copy_code_button)
        cl.addLayout(options)
        self.code_preview = QPlainTextEdit()
        self.code_preview.setObjectName('codeArea')
        self.code_preview.setReadOnly(True)
        self.code_preview.setPlaceholderText('Select one or more rows to preview framework bindings.')
        cl.addWidget(self.code_preview, 1)
        lower_split.addTab(code_panel, 'Code Preview')
        right.addWidget(lower_split)
        right.setSizes([570, 210])
        self.switch_mode(0)

    def build_exclusion_menu(self):
        self.exclusion_actions = []
        self.exclusion_menu = PersistentCheckMenu(self)
        self.exclusion_menu.setMinimumWidth(290)
        type_heading = QAction('LOCATOR TYPES', self.exclusion_menu)
        type_heading.setEnabled(False)
        type_font = type_heading.font()
        type_font.setBold(True)
        type_heading.setFont(type_font)
        self.exclusion_menu.addAction(type_heading)
        for title, option_type, value in [('CSS selectors', 'type', 'CSS'), ('XPath selectors', 'type', 'XPath')]:
            action = QAction(title, self.exclusion_menu)
            action.setCheckable(True)
            action.setData((option_type, value))
            action.toggled.connect(self.update_exclusion_label)
            self.exclusion_menu.addAction(action)
            self.exclusion_actions.append((action, option_type, value, title))
        self.exclusion_menu.addSeparator()
        strategy_heading = QAction('LOCATOR STRATEGIES', self.exclusion_menu)
        strategy_heading.setEnabled(False)
        strategy_font = strategy_heading.font()
        strategy_font.setBold(True)
        strategy_heading.setFont(strategy_font)
        self.exclusion_menu.addAction(strategy_heading)
        for strategy in LOCATOR_STRATEGIES:
            action = QAction(strategy, self.exclusion_menu)
            action.setCheckable(True)
            action.setData(('strategy', strategy))
            action.toggled.connect(self.update_exclusion_label)
            self.exclusion_menu.addAction(action)
            self.exclusion_actions.append((action, 'strategy', strategy, strategy))
        self.exclusion_menu.addSeparator()
        clear_exclusions_action = QAction('Clear selections', self.exclusion_menu)
        self.exclusion_menu.addAction(clear_exclusions_action)
        self.exclusion_menu.clear_requested.connect(self.clear_exclusions)

    def build_picker_workspace(self):
        outer = self.workspace('Element Picker', 'Browse inside the app, click elements to capture them, and keep each generated locator bundle separate.')
        toolbar_panel, toolbar_layout = panel('IN-APP BROWSER')
        bar = QHBoxLayout()
        self.picker_url_entry = QLineEdit('https://example.com')
        self.picker_url_entry.setPlaceholderText('Enter a page URL')
        self.picker_url_entry.returnPressed.connect(self.navigate_picker)
        bar.addWidget(self.picker_url_entry, 1)
        self.picker_go_button = QPushButton('Go')
        self.picker_go_button.setObjectName('primaryButton')
        self.picker_go_button.clicked.connect(self.navigate_picker)
        bar.addWidget(self.picker_go_button)
        self.picker_close_page_button = QPushButton('Close page')
        self.picker_close_page_button.setToolTip('Close the browser page and give the picker panels more room.')
        self.picker_close_page_button.setEnabled(False)
        self.picker_close_page_button.clicked.connect(self.close_picker_page)
        bar.addWidget(self.picker_close_page_button)
        self.picker_scan_button = QPushButton('Scan whole page')
        self.picker_scan_button.setObjectName('primaryButton')
        self.picker_scan_button.setToolTip('Generate locators for eligible elements across the URL in the picker address bar.')
        self.picker_scan_button.setEnabled(HAS_SELENIUM or (HAS_WEBENGINE and self._picker_loaded))
        self.picker_scan_button.clicked.connect(self.scan_picker_whole_page)
        bar.addWidget(self.picker_scan_button)
        browser_actions = QHBoxLayout()
        self.picker_back_button = QPushButton('Back')
        self.picker_back_button.setEnabled(False)
        self.picker_back_button.clicked.connect(lambda: self.picker_view.back() if self.picker_view else None)
        browser_actions.addWidget(self.picker_back_button)
        self.picker_forward_button = QPushButton('Forward')
        self.picker_forward_button.setEnabled(False)
        self.picker_forward_button.clicked.connect(lambda: self.picker_view.forward() if self.picker_view else None)
        browser_actions.addWidget(self.picker_forward_button)
        self.picker_reload_button = QPushButton('Reload')
        self.picker_reload_button.setEnabled(False)
        self.picker_reload_button.clicked.connect(lambda: self.picker_view.reload() if self.picker_view else None)
        browser_actions.addWidget(self.picker_reload_button)
        self.picker_pick_mode = QCheckBox('Click to capture')
        self.picker_pick_mode.setChecked(True)
        self.picker_pick_mode.setEnabled(False)
        self.picker_pick_mode.toggled.connect(self.update_pick_mode)
        browser_actions.addWidget(self.picker_pick_mode)
        self.toggle_picker_panels = QToolButton()
        self.toggle_picker_panels.setText('Hide details  ◀')
        self.toggle_picker_panels.setToolTip('Collapse or reopen the captured elements, candidate locators, and code preview panel.')
        self.toggle_picker_panels.setCheckable(True)
        self.toggle_picker_panels.setChecked(True)
        self.toggle_picker_panels.toggled.connect(self.toggle_picker_details)
        browser_actions.addWidget(self.toggle_picker_panels)
        browser_actions.addStretch(1)
        toolbar_layout.addLayout(bar)
        toolbar_layout.addLayout(browser_actions)
        picker_ignore_row = QHBoxLayout()
        self.picker_exclude_button = ArrowMenuButton()
        self.picker_exclude_button.setObjectName('excludeButton')
        self.picker_exclude_button.setText('Ignore patterns')
        self.picker_exclude_button.setToolTip('Choose locator types or strategies to omit from future page scans and captured-element generation.')
        self.picker_exclude_button.setPopupMode(QToolButton.ToolButtonPopupMode.InstantPopup)
        self.picker_exclude_button.setMenu(self.exclusion_menu)
        picker_ignore_row.addWidget(self.picker_exclude_button)
        picker_ignore_row.addWidget(label('Framework'))
        self.picker_framework_combo = ArrowComboBox()
        self.picker_framework_combo.addItems(FRAMEWORKS)
        self.picker_framework_combo.setCurrentText(self.framework)
        self.picker_framework_combo.currentTextChanged.connect(self.change_framework)
        picker_ignore_row.addWidget(self.picker_framework_combo)
        self.picker_excluded_summary = label('Ignored: none', True)
        self.picker_excluded_summary.setWordWrap(True)
        picker_ignore_row.addWidget(self.picker_excluded_summary, 1)
        toolbar_layout.addLayout(picker_ignore_row)
        self.picker_status_label = label(
            'Website closed. Generate locators from HTML below, or enter a URL and click Go.' if HAS_WEBENGINE else
            'Embedded browser unavailable; the HTML locator generator below is ready to use.', True)
        toolbar_layout.addWidget(self.picker_status_label)
        outer.addWidget(toolbar_panel)

        split = QSplitter(Qt.Orientation.Horizontal)
        self.picker_splitter = split
        self.register_responsive_splitter(split, 1080, ratio=0.56, vertical_ratio=0.58)
        outer.addWidget(split, 1)
        browser_panel, browser_layout = panel('PAGE')
        self.picker_display_title = browser_layout.itemAt(0).widget()
        self.picker_view = QWebEngineView() if HAS_WEBENGINE else None
        self.picker_page_stack = QStackedWidget()
        fallback_page = QWidget()
        fallback_layout = QVBoxLayout(fallback_page)
        fallback_layout.setContentsMargins(10, 10, 10, 10)
        fallback_layout.setSpacing(8)
        fallback_layout.addWidget(label('Generate locators from pasted HTML', True))
        self.picker_generator_html = QPlainTextEdit()
        self.picker_generator_html.setObjectName('codeArea')
        self.picker_generator_html.setPlaceholderText('Paste an HTML element or fragment here.')
        self.picker_generator_html.setPlainText(EXAMPLE_HTML)
        fallback_layout.addWidget(self.picker_generator_html, 1)
        fallback_actions = QHBoxLayout()
        self.picker_generate_button = QPushButton('Generate Locators')
        self.picker_generate_button.setObjectName('primaryButton')
        self.picker_generate_button.clicked.connect(self.generate_picker_fallback)
        fallback_actions.addWidget(self.picker_generate_button)
        picker_example_button = QPushButton('Generate example')
        picker_example_button.clicked.connect(lambda: self.picker_generator_html.setPlainText(EXAMPLE_HTML))
        fallback_actions.addWidget(picker_example_button)
        picker_fallback_copy = QPushButton('Copy selectors')
        picker_fallback_copy.clicked.connect(self.copy_picker_candidate_selectors)
        fallback_actions.addWidget(picker_fallback_copy)
        fallback_layout.addLayout(fallback_actions)
        self.picker_fallback_page = fallback_page
        self.picker_page_stack.addWidget(fallback_page)
        if self.picker_view:
            self.picker_view.setMinimumSize(280, 190)
            self.picker_view.urlChanged.connect(self.on_picker_url_changed)
            self.picker_view.loadFinished.connect(self.on_picker_load_finished)
            self.picker_page_stack.addWidget(self.picker_view)
        else:
            unavailable = label('This build needs the PySide6 Qt WebEngine component for the embedded picker.', True)
            unavailable.setWordWrap(True)
            self.picker_page_stack.addWidget(unavailable)
        browser_layout.addWidget(self.picker_page_stack, 1)
        split.addWidget(browser_panel)
        self.picker_browser_panel = browser_panel

        picker_side = QTabWidget()
        picker_side.setDocumentMode(True)
        picker_side.setMinimumWidth(300)
        self.picker_detail_tabs = picker_side
        self.picker_detail_container = QWidget()
        picker_detail_layout = QVBoxLayout(self.picker_detail_container)
        picker_detail_layout.setContentsMargins(0, 0, 0, 0)
        picker_detail_layout.setSpacing(7)
        picker_detail_layout.addWidget(picker_side, 1)
        bundles_panel, bundles_layout = panel('CAPTURED ELEMENT BUNDLES')
        bundles_help = label('Check one or more captured bundles. Candidates and Code Preview combine the checked bundles and refresh as you select or clear them. Copy all and Copy code use the same checked bundles; expand a bundle to inspect individual selectors.', True)
        bundles_help.setWordWrap(True)
        bundles_layout.addWidget(bundles_help)
        capture_selection_row = QHBoxLayout()
        capture_selection_row.addWidget(label('Selection mode'))
        self.capture_selection_mode = ArrowComboBox()
        self.capture_selection_mode.addItems(['Multiple selection', 'Single selection'])
        self.capture_selection_mode.currentTextChanged.connect(self.set_capture_selection_mode)
        capture_selection_row.addWidget(self.capture_selection_mode)
        self.select_all_captures = QCheckBox('Select all')
        self.select_all_captures.toggled.connect(self.set_all_captures_checked)
        capture_selection_row.addWidget(self.select_all_captures)
        capture_selection_row.addStretch(1)
        bundles_layout.addLayout(capture_selection_row)
        self.captured_bundle_list = QTreeWidget()
        self.captured_bundle_list.setHeaderLabels(['Captured element / locator', 'Type / count', 'Strategy'])
        self.captured_bundle_list.setRootIsDecorated(True)
        self.captured_bundle_list.setAlternatingRowColors(True)
        self.captured_bundle_list.setUniformRowHeights(True)
        self.captured_bundle_list.setEditTriggers(QAbstractItemView.EditTrigger.NoEditTriggers)
        self.captured_bundle_list.setColumnWidth(0, 260)
        self.captured_bundle_list.setColumnWidth(1, 95)
        self.captured_bundle_list.currentItemChanged.connect(self.update_open_bundle_button)
        self.captured_bundle_list.currentItemChanged.connect(self.update_picker_candidate_view)
        self.captured_bundle_list.itemChanged.connect(self.on_capture_bundle_check_changed)
        self.captured_bundle_list.itemClicked.connect(self.on_capture_tree_clicked)
        self.captured_bundle_list.itemDoubleClicked.connect(self.on_capture_tree_double_clicked)
        bundles_layout.addWidget(self.captured_bundle_list, 1)
        copy_actions = QHBoxLayout()
        self.copy_capture_locator_button = QPushButton('Copy one')
        self.copy_capture_locator_button.setToolTip('Copy the selected locator expression.')
        self.copy_capture_locator_button.setEnabled(False)
        self.copy_capture_locator_button.clicked.connect(self.copy_selected_capture_locator)
        copy_actions.addWidget(self.copy_capture_locator_button)
        self.copy_capture_bundle_button = QPushButton('Copy all')
        self.copy_capture_bundle_button.setToolTip('Copy every locator expression from the checked bundles, or the current bundle when none are checked.')
        self.copy_capture_bundle_button.setEnabled(False)
        self.copy_capture_bundle_button.clicked.connect(self.copy_capture_bundle_selectors)
        copy_actions.addWidget(self.copy_capture_bundle_button)
        self.copy_capture_code_button = QPushButton('Copy code')
        self.copy_capture_code_button.setToolTip('Copy framework bindings for every locator in the selected bundle.')
        self.copy_capture_code_button.setEnabled(False)
        self.copy_capture_code_button.clicked.connect(self.copy_capture_bundle_code)
        copy_actions.addWidget(self.copy_capture_code_button)
        bundles_layout.addLayout(copy_actions)
        bundle_actions = QHBoxLayout()
        self.open_capture_button = QPushButton('Open locators')
        self.open_capture_button.setObjectName('primaryButton')
        self.open_capture_button.setEnabled(False)
        self.open_capture_button.clicked.connect(self.open_selected_capture)
        bundle_actions.addWidget(self.open_capture_button)
        clear_bundles = QPushButton('Clear bundles')
        clear_bundles.clicked.connect(self.clear_captured_bundles)
        bundle_actions.addWidget(clear_bundles)
        bundles_layout.addLayout(bundle_actions)
        picker_side.addTab(bundles_panel, 'Captured')

        candidate_panel, candidate_layout = panel('CANDIDATE LOCATORS')
        self.picker_candidates_caption = label('Select one or more captured bundles to view their candidates.', True)
        candidate_layout.addWidget(self.picker_candidates_caption)
        picker_selection_row = QHBoxLayout()
        picker_selection_row.addWidget(label('Selection mode'))
        self.picker_selection_mode = ArrowComboBox()
        self.picker_selection_mode.addItems(['Multiple selection', 'Single selection'])
        self.picker_selection_mode.currentTextChanged.connect(self.set_picker_selection_mode)
        picker_selection_row.addWidget(self.picker_selection_mode)
        picker_selection_row.addStretch(1)
        candidate_layout.addLayout(picker_selection_row)
        self.picker_candidate_tree = table(['Select', 'Selector', 'Type', 'Strategy', 'Score', 'Matches', 'Captured element'],
                                           [48, 250, 65, 135, 90, 75, 190])
        self.picker_candidate_tree.itemChanged.connect(self.on_picker_candidate_changed)
        self.picker_candidate_tree.itemSelectionChanged.connect(self.update_picker_locator_details)
        candidate_layout.addWidget(self.picker_candidate_tree, 1)
        self.picker_candidate_sources = []
        self.picker_select_all = QCheckBox('Select all locators')
        self.picker_select_all.toggled.connect(self.set_all_picker_candidates)
        candidate_layout.addWidget(self.picker_select_all)
        picker_side.addTab(candidate_panel, 'Candidates')

        picker_code_panel, picker_code_layout = panel('CODE PREVIEW')
        picker_code_options = QHBoxLayout()
        picker_code_options.addWidget(label('Selected locator bindings'))
        picker_code_options.addStretch(1)
        picker_code_layout.addLayout(picker_code_options)
        self.picker_code_preview = QPlainTextEdit()
        self.picker_code_preview.setObjectName('codeArea')
        self.picker_code_preview.setReadOnly(True)
        self.picker_code_preview.setPlaceholderText('Select one or more candidate locators to preview framework code.')
        picker_code_layout.addWidget(self.picker_code_preview, 1)
        copy_picker_preview = QPushButton('Copy preview code')
        copy_picker_preview.setObjectName('primaryButton')
        copy_picker_preview.clicked.connect(self.copy_picker_candidate_code)
        picker_code_layout.addWidget(copy_picker_preview)
        picker_side.addTab(picker_code_panel, 'Code Preview')

        locator_details_panel, locator_details_layout = panel('LOCATOR DETAILS')
        self.picker_locator_details_panel = locator_details_panel
        self.picker_locator_details = QPlainTextEdit()
        self.picker_locator_details.setReadOnly(True)
        self.picker_locator_details.setObjectName('codeArea')
        self.picker_locator_details.setPlaceholderText('Select a captured element or candidate locator to inspect its details.')
        locator_details_layout.addWidget(self.picker_locator_details, 1)
        picker_side.addTab(locator_details_panel, 'Locator Details')
        self.picker_action_widget = QWidget()
        picker_action_row = QGridLayout(self.picker_action_widget)
        self.picker_copy_selectors_button = QPushButton('Copy selectors')
        self.picker_copy_selectors_button.clicked.connect(self.copy_picker_candidate_selectors)
        picker_action_row.addWidget(self.picker_copy_selectors_button, 0, 0)
        picker_add_pom_button = QPushButton('Add to POM')
        picker_add_pom_button.clicked.connect(self.add_picker_candidates_to_pom)
        picker_action_row.addWidget(picker_add_pom_button, 0, 1)
        picker_send_audit_button = QPushButton('Send to Health Audit')
        picker_send_audit_button.clicked.connect(self.send_picker_candidates_to_audit)
        picker_action_row.addWidget(picker_send_audit_button, 0, 2)
        self.picker_test_button = QPushButton('Test live')
        self.picker_test_button.setEnabled(HAS_SELENIUM)
        self.picker_test_button.clicked.connect(self.test_picker_candidates_live)
        picker_action_row.addWidget(self.picker_test_button, 1, 0)
        self.picker_export_button = QPushButton('Export')
        self.picker_export_button.clicked.connect(self.export_picker_candidates)
        picker_action_row.addWidget(self.picker_export_button, 1, 1)
        picker_copy_code_button = QPushButton('Copy code')
        picker_copy_code_button.clicked.connect(self.copy_picker_candidate_code)
        picker_action_row.addWidget(picker_copy_code_button, 1, 2)
        picker_detail_layout.addWidget(self.picker_action_widget, 0)
        split.addWidget(self.picker_detail_container)
        split.setSizes([660, 510])

        if self.picker_view:
            self.picker_poll_timer = QTimer(self)
            self.picker_poll_timer.setInterval(250)
            self.picker_poll_timer.timeout.connect(self.poll_picker_selection)
        self.picker_page_stack.setCurrentWidget(self.picker_fallback_page)

    def build_pom_workspace(self):
        outer = self.workspace('Page Objects', 'Collect selected locators and generate class code using your preferred framework.')
        bar, bl = panel('PAGE OBJECT SETTINGS')
        controls = QHBoxLayout()
        controls.addWidget(label('Class name'))
        self.pom_class_entry = QLineEdit('LoginPage')
        self.pom_class_entry.setMaximumWidth(235)
        self.pom_class_entry.textChanged.connect(self.update_pom_code_preview)
        controls.addWidget(self.pom_class_entry)
        controls.addWidget(label('Framework'))
        self.pom_framework = ArrowComboBox()
        self.pom_framework.addItems(FRAMEWORKS)
        self.pom_framework.currentTextChanged.connect(self.change_framework)
        controls.addWidget(self.pom_framework)
        controls.addStretch(1)
        clear = QPushButton('Clear elements')
        clear.clicked.connect(self.clear_pom)
        controls.addWidget(clear)
        bl.addLayout(controls)
        outer.addWidget(bar)

        split = QSplitter(Qt.Orientation.Horizontal)
        self.register_responsive_splitter(split, 1180)
        outer.addWidget(split, 1)
        elems, el = panel('SELECTED ELEMENTS')
        self.pom_count = label('No elements added. Select locators from Element Picker.', True)
        el.addWidget(self.pom_count)
        self.pom_tree = table(['Variable', 'Type', 'Selector'], [145, 72, 300])
        self.pom_tree.itemSelectionChanged.connect(self.load_selected_pom_variable)
        el.addWidget(self.pom_tree, 1)
        variable_row = QHBoxLayout()
        variable_row.addWidget(label('Variable name'))
        self.pom_variable_entry = QLineEdit()
        self.pom_variable_entry.setPlaceholderText('Select an element, then enter a variable name')
        self.pom_variable_entry.returnPressed.connect(self.rename_selected_pom_variable)
        variable_row.addWidget(self.pom_variable_entry, 1)
        self.rename_pom_button = QPushButton('Set name')
        self.rename_pom_button.clicked.connect(self.rename_selected_pom_variable)
        variable_row.addWidget(self.rename_pom_button)
        el.addLayout(variable_row)
        split.addWidget(elems)
        preview_panel, pl = panel('GENERATED CLASS CODE')
        self.pom_code_text = QPlainTextEdit()
        self.pom_code_text.setObjectName('codeArea')
        self.pom_code_text.setReadOnly(True)
        pl.addWidget(self.pom_code_text, 1)
        self.copy_pom_button = QPushButton('Copy POM class code')
        self.copy_pom_button.setObjectName('primaryButton')
        self.copy_pom_button.clicked.connect(self.copy_pom_class)
        pl.addWidget(self.copy_pom_button)
        split.addWidget(preview_panel)
        split.setSizes([430, 710])
        self.update_pom_code_preview()

    def toggle_picker_details(self, visible):
        if not hasattr(self, 'picker_detail_container'):
            return
        # Hide the complete right pane, including the action bar, so the page
        # or HTML editor can take over the available splitter space.
        self.picker_detail_container.setVisible(visible)
        self.picker_detail_tabs.setVisible(visible)
        self.picker_action_widget.setVisible(visible)
        self.toggle_picker_panels.setText('Hide details  ◀' if visible else 'Show details  ▶')
        extent = (self.picker_splitter.height()
                  if self.picker_splitter.orientation() == Qt.Orientation.Vertical
                  else self.picker_splitter.width())
        if not visible:
            self.picker_splitter.setSizes([max(1, extent), 0])
            return
        first_ratio = (0.58 if self.picker_splitter.orientation() == Qt.Orientation.Horizontal
                       else 0.50)
        first = max(1, round(extent * first_ratio))
        self.picker_splitter.setSizes([first, max(1, extent - first)])

    def build_audit_workspace(self):
        outer = self.workspace('Health Audit', 'Check selectors against a live page; review tentative fallback suggestions separately.')
        ctrl, cl = panel('AUDIT TARGET')
        bar = QHBoxLayout()
        bar.addWidget(label('Target URL'))
        self.audit_url_entry = QLineEdit('https://example.com')
        bar.addWidget(self.audit_url_entry, 1)
        self.audit_button = QPushButton('Run Health Audit')
        self.audit_button.setObjectName('primaryButton')
        self.audit_button.clicked.connect(self.run_health_audit_and_heal)
        self.audit_button.setEnabled(HAS_SELENIUM)
        bar.addWidget(self.audit_button)
        cl.addLayout(bar)
        outer.addWidget(ctrl)
        main_split = QSplitter(Qt.Orientation.Horizontal)
        self.register_responsive_splitter(main_split, 1180)
        outer.addWidget(main_split, 1)
        inputs, il = panel('LOCATORS TO AUDIT')
        il.addWidget(label('Paste selectors or Selenium locator declarations, one per line.', True))
        self.audit_input_text = QPlainTextEdit()
        self.audit_input_text.setObjectName('codeArea')
        self.audit_input_text.setPlaceholderText(
            'CSS/XPath selectors or Selenium declarations:\n'
            '//h1\n#submit-button\nBy loginButton = By.id("login")\n'
            'submit = (By.CSS_SELECTOR, "button[type=submit]")'
        )
        il.addWidget(self.audit_input_text, 1)
        main_split.addWidget(inputs)
        rhs = QSplitter(Qt.Orientation.Vertical)
        rhs.setChildrenCollapsible(False)
        result, rl = panel('LIVE VALIDATION RESULTS')
        rl.addWidget(label('Only the original locators are checked with Selenium.', True))
        self.audit_tree = table(['Locator', 'Matches', 'Status'], [390, 90, 180])
        rl.addWidget(self.audit_tree, 1)
        rhs.addWidget(result)
        fallbacks, fl = panel('UNVERIFIED FALLBACK SUGGESTIONS')
        desc = label('Ranked alternatives based on matching tag, attributes, classes, or text in the captured page snapshot. Suggestions are not live-validated.', True)
        desc.setWordWrap(True)
        fl.addWidget(desc)
        self.suggestion_tree = table(['Original', 'Suggested alternative', 'Notes'], [285, 300, 270])
        fl.addWidget(self.suggestion_tree, 1)
        rhs.addWidget(fallbacks)
        rhs.setSizes([280, 240])
        main_split.addWidget(rhs)
        main_split.setSizes([350, 800])

    def build_help_workspace(self):
        outer = self.workspace('Help & Guide', 'A practical guide to generating, selecting, validating, and exporting locators.')
        contents, body = panel('LOCATOR GENERATOR  ·  USER GUIDE')
        guide = QTextBrowser()
        guide.setReadOnly(True)
        guide.setOpenExternalLinks(False)
        guide.setObjectName('guideBrowser')
        self.help_browser = guide
        guide.setHtml('''
<html><body style="font-family:Segoe UI,Arial,sans-serif; font-size:12pt;">
<h2 style="color:#3b82f6; margin-bottom:4px;">Start here</h2>
<p>Use Element Picker to open an embedded page, capture elements, scan the whole page, or generate ranked CSS and XPath options from pasted HTML.</p>
<p><b>Install:</b> <code>pip install PySide6 beautifulsoup4 selenium openpyxl lxml</code><br>
<b>Run:</b> <code>python locator_generator_pyside6.py</code></p>
<hr>
<h3 style="color:#3b82f6;">1 · Element Picker</h3>
<ul>
 <li><b>Generate from HTML:</b> When the embedded page is closed, paste an element or fragment into the HTML editor. The app chooses a useful target and generates ranked locator candidates.</li>
 <li><b>Element Picker:</b> Use Go to open the embedded page, or Scan whole page to add page-wide candidates as a bundle. Close page reveals the HTML editor, where you can generate locators, load an example, or copy selectors. Ignore patterns apply to new scans and captured-element generation. Select captured bundles singly or in groups; Candidate Locators and Code Preview update from checked bundles. Clicking a captured bundle or selector does not change the open details tab. The persistent action bar includes Copy selectors, Add to POM, Send to Health Audit, Test live, and Export.</li>
 <li><b>Scan Page:</b> Enter a URL to inspect buttons, inputs, links, forms, and elements with test IDs. Scan results come from a page-source snapshot.</li>
 <li><b>Select locators:</b> Tick the checkbox beside each selector, or use <b>Select all generated selectors</b>. Copy, add to Page Objects, send to Health Audit, and Test live use the checked selectors.</li>
 <li><b>Inspect and copy:</b> Open Locator Details to read the selected locator's strategy, score, snapshot match count, notes, and framework binding. It says <i>Live test: Not tested yet</i> until you run Test live; afterward it shows the actual match count, no-match result, or evaluation error for the tested URL. Choose a framework beside Ignore patterns to preview or copy its code binding.</li>
 <li><b>Ignore patterns:</b> Choose CSS, XPath, or one or more locator strategies. The grouped checklist stays open while you choose multiple items; active exclusions remain visible above the results and apply to HTML generation and page scans.</li>
</ul>
<p><i>Fragment and scan counts are calculated from captured HTML. They do not prove that a selector works on the live page.</i></p>
<hr>
<h3 style="color:#3b82f6;">2 · Page Objects</h3>
<ul>
 <li>Check locators in Element Picker, then choose <b>Add to POM</b>. Duplicate expressions are skipped and new variable names are made unique automatically.</li>
 <li>On the Page Objects page, select one element, enter a variable name, and click <b>Set name</b> (or press Enter). Names must start with a letter or underscore and contain only letters, digits, or underscores; names already in use are rejected.</li>
 <li>Edit the class name to refresh the code. The framework selector is a bordered dropdown with an arrow; changing it updates both code previews.</li>
 <li>Use <b>Clear elements</b> to empty the Page Object list, or <b>Copy POM class code</b> to copy the generated class.</li>
</ul>
<hr>
<h3 style="color:#3b82f6;">3 · Health Audit</h3>
<ul>
 <li>Enter a page URL and one CSS or XPath locator per line, then run the audit. Selenium checks each original locator against the live page.</li>
 <li><b>Found</b>, <b>0 matches</b>, and <b>Evaluation error</b> are different results. An evaluation error means Selenium could not evaluate the selector.</li>
 <li>For zero-match locators, the app may show up to three ranked alternatives based on matching tags, IDs, classes, attributes, or text in the captured page snapshot. Notes explain the match evidence and snapshot count.</li>
 <li>Fallbacks are suggestions only. They are not tested against the live page and do not replace or repair the original selector.</li>
</ul>
<hr>
<h3 style="color:#3b82f6;">4 · Export</h3>
<ul>
 <li><b>Excel (.xlsx):</b> Exports the generated locator list with a clean white cell background, readable text, wrapped long expressions, filters, and a frozen header row.</li>
 <li><b>CSV (.csv):</b> Exports the same columns as plain data. CSV does not store colors, filters, or column formatting.</li>
 <li>Exports include locator type and strategy, expression, stability rank and score, match count, framework binding, notes, and match source.</li>
</ul>
<hr>
<h3 style="color:#3b82f6;">Frameworks and display</h3>
<p>The framework dropdown supports <b>Java Selenium</b>, <b>Python Selenium</b>, <b>Playwright</b>, and <b>Cypress</b>. Use the header button to switch between dark and light app themes. The Help page follows the active app theme.</p>
<h3 style="color:#3b82f6;">Keyboard and dependencies</h3>
<ul>
 <li>Use the header menu button to open or close the workspace drawer and give the active page more room.</li>
 <li><b>Ctrl+1 … Ctrl+4:</b> Switch between Element Picker, Page Objects, Health Audit, and Help. Zoom can go down to 50% and up to 150%.</li>
 <li><b>Zoom:</b> Use the A− and A+ buttons or Ctrl+- / Ctrl++ to resize the interface; Ctrl+0 restores 100%. Split views and navigation adapt to narrower windows.</li>
 <li><b>Tab / Shift+Tab:</b> Move between controls.</li>
 <li><b>PySide6 with Qt WebEngine:</b> Required for the desktop app and embedded Element Picker. <b>BeautifulSoup:</b> Required for HTML parsing. <b>Selenium and Chrome or Edge:</b> Required for live scans, live tests, and Health Audit; Selenium Manager may need internet access to set up the matching driver. <b>openpyxl:</b> Required for Excel export; CSV works without it. <b>lxml:</b> Improves XPath match counts.</li>
</ul>
<p><i>Scores and fallback rankings are heuristics, not guarantees. Some framework snippets may need project-specific imports or syntax adjustments.</i></p>
</body></html>
''')
        body.addWidget(guide, 1)
        copy = QPushButton('Copy help text')
        copy.clicked.connect(lambda: self.copy(guide.toPlainText(), 'Help copied.'))
        body.addWidget(copy)
        outer.addWidget(contents, 1)

    def switch_workspace(self, index):
        self.workspaces.setCurrentIndex(index)
        for number, button in enumerate(self.nav_buttons):
            button.setChecked(number == index)
        if hasattr(self, 'responsive_nav') and self.responsive_nav.currentIndex() != index:
            self.responsive_nav.blockSignals(True)
            self.responsive_nav.setCurrentIndex(index)
            self.responsive_nav.blockSignals(False)

    def switch_mode(self, index):
        self.modes.setCurrentIndex(index)
        for number, button in enumerate(self.mode_buttons):
            button.setChecked(number == index)
        self.generate_button.setVisible(index == 0)
        self.scan_button.setEnabled(not self.busy and HAS_SELENIUM)
        if not HAS_SELENIUM and index == 1:
            self.set_status('Selenium is required for live page scans. The embedded picker does not need Selenium.')

    def set_status(self, message):
        self.status.setText(message)

    def popup(self, title, message):
        QMessageBox.warning(self, title, message)

    def get_excluded_locator_options(self):
        excluded_types, excluded_strategies = [], []
        for action, option_type, value, _title in self.exclusion_actions:
            if not action.isChecked():
                continue
            (excluded_types if option_type == 'type' else excluded_strategies).append(value)
        return excluded_types, excluded_strategies

    def update_exclusion_label(self, *_):
        selected = [title for action, _kind, _value, title in self.exclusion_actions
                    if action.isChecked()]
        count = len(selected)
        label_text = f'Ignore patterns ({count})' if count else 'Ignore patterns'
        if hasattr(self, 'exclude_button'):
            self.exclude_button.setText(label_text)
        if hasattr(self, 'excluded_summary'):
            self.excluded_summary.setText('Ignored: ' + (', '.join(selected) if selected else 'none'))
        if hasattr(self, 'picker_exclude_button'):
            self.picker_exclude_button.setText(label_text)
            self.picker_excluded_summary.setText('Ignored: ' + (', '.join(selected) if selected else 'none'))

    def clear_exclusions(self):
        for action, _kind, _value, _title in self.exclusion_actions:
            action.setChecked(False)
        self.update_exclusion_label()

    def generate(self):
        if BeautifulSoup is None:
            self.popup('Dependency missing', 'Install BeautifulSoup: pip install beautifulsoup4')
            return
        markup = self.html.toPlainText().strip()
        if not markup:
            self.popup('HTML required', 'Paste HTML before generating locators.')
            self.html.setFocus()
            return
        self.loaded_bundle_id = None
        excluded_types, excluded_strategies = self.get_excluded_locator_options()
        self.launch_job(
            lambda: generate_locators_job(markup, excluded_types, excluded_strategies),
            self.populate_generated_results,
            'Generating locator candidates…',
        )

    def populate_generated_results(self, payload):
        self.rows = payload['rows']
        self.count_context = 'HTML fragment'
        self.show_candidates(f'Target <{payload["tag"]}> · {len(self.rows)} candidates')

    def show_candidates(self, summary):
        self.put_candidate_rows()
        self.results_caption.setText(summary + f'   |   Match source: {self.count_context}')
        self.preview()
        self.set_status(summary)

    def clear_html(self):
        self.html.clear()
        self.rows = []
        put_rows(self.tree, [])
        self.select_all_candidates.setChecked(False)
        self.results_caption.setText('No locators generated yet.')
        self.preview()
        self.set_status('HTML and locator results cleared.')

    def selected_indices(self):
        return sorted({index.row() for index in self.tree.selectedIndexes()})

    def selected_candidates(self):
        return [self.rows[i] for i in range(min(len(self.rows), self.tree.rowCount()))
                if self.tree.item(i, 0) is not None
                and self.tree.item(i, 0).checkState() == Qt.CheckState.Checked]

    def put_candidate_rows(self):
        self.tree.blockSignals(True)
        try:
            self.tree.setRowCount(0)
            for candidate in self.rows:
                row = self.tree.rowCount()
                self.tree.insertRow(row)
                check = QTableWidgetItem('')
                check.setFlags(check.flags() | Qt.ItemFlag.ItemIsUserCheckable | Qt.ItemFlag.ItemIsEnabled)
                check.setCheckState(Qt.CheckState.Unchecked)
                self.tree.setItem(row, 0, check)
                values = (candidate['expr'], candidate['kind'], candidate['strategy'],
                          f'{candidate["rank"]} · {candidate["score"]}/100',
                          str(candidate['count']) if candidate['count'] is not None else 'N/A')
                for column, value in enumerate(values, 1):
                    self.tree.setItem(row, column, QTableWidgetItem(str(value)))
        finally:
            self.tree.blockSignals(False)
        self.select_all_candidates.blockSignals(True)
        self.select_all_candidates.setChecked(False)
        self.select_all_candidates.blockSignals(False)

    def set_all_candidates_checked(self, checked):
        self.tree.blockSignals(True)
        try:
            for row in range(self.tree.rowCount()):
                item = self.tree.item(row, 0)
                if item is not None:
                    item.setCheckState(Qt.CheckState.Checked if checked else Qt.CheckState.Unchecked)
        finally:
            self.tree.blockSignals(False)
        self.preview()

    def on_candidate_check_changed(self, item):
        if item.column() != 0:
            return
        all_checked = bool(self.rows) and all(
            self.tree.item(row, 0) is not None
            and self.tree.item(row, 0).checkState() == Qt.CheckState.Checked
            for row in range(min(len(self.rows), self.tree.rowCount()))
        )
        self.select_all_candidates.blockSignals(True)
        self.select_all_candidates.setChecked(all_checked)
        self.select_all_candidates.blockSignals(False)
        self.preview()

    def preview(self):
        candidates = self.selected_candidates()
        if not candidates:
            self.selected_note.setPlainText('Select a locator to inspect its strategy and notes.')
            self.code_preview.setPlainText('Select one or more rows to preview framework code.')
            return
        c = candidates[0]
        count = c['count'] if c['count'] is not None else 'N/A'
        self.selected_note.setPlainText(
            f'Locator: {c["expr"]}\n'
            f'Type: {c["kind"]}   Strategy: {c["strategy"]}\n'
            f'Score: {c["score"]}/100 ({c["rank"]})\n'
            f'{self.count_context} matches: {count}\n\n'
            f'{c["note"]}\n\n'
            'These counts do not establish live Selenium verification.'
        )
        self.code_preview.setPlainText('\n'.join(format_framework_code(item, self.framework) for item in candidates))

    def change_framework(self, value):
        if value not in FRAMEWORKS:
            return
        self.framework = value
        for combo in (getattr(self, 'framework_combo', None), getattr(self, 'pom_framework', None), getattr(self, 'picker_framework_combo', None)):
            if combo is not None and combo.currentText() != value:
                was_blocked = combo.blockSignals(True)
                combo.setCurrentText(value)
                combo.blockSignals(was_blocked)
        if hasattr(self, 'code_preview'):
            self.preview()
        if hasattr(self, 'picker_code_preview'):
            self.update_picker_code_preview()
            self.update_picker_locator_details()
        self.update_pom_code_preview()

    def copy(self, content, status):
        QApplication.clipboard().setText(content)
        self.set_status(status)

    def copy_expr(self):
        rows = self.selected_candidates()
        if not rows:
            self.set_status('Select one or more locator rows to copy.')
            return
        self.copy('\n'.join(c['expr'] for c in rows), f'Copied {len(rows)} selector(s).')

    def copy_framework_code(self):
        rows = self.selected_candidates()
        if not rows:
            self.set_status('Select one or more locator rows to copy code.')
            return
        self.copy('\n'.join(format_framework_code(c, self.framework) for c in rows),
                  f'Copied {len(rows)} {self.framework} binding(s).')

    def add_selected_to_pom(self):
        candidates = self.selected_candidates()
        if not candidates:
            self.set_status('Select one or more rows to add to Page Objects.')
            return
        added = duplicates = 0
        for c in candidates:
            expr = c['expr']
            if any(e['expr'] == expr for e in self.pom_elements):
                duplicates += 1
                continue
            original = generate_descriptive_name(expr)
            if self.loaded_bundle_id is not None:
                bundle = next((entry for entry in self.captured_bundles
                               if entry['id'] == self.loaded_bundle_id), None)
                if bundle and bundle.get('variable_prefix'):
                    original = f'{bundle["variable_prefix"]}_{original}'
            variable = original
            counter = 1
            while variable in {e['varname'] for e in self.pom_elements}:
                variable = f'{original}{counter}'
                counter += 1
            self.pom_elements.append({'varname': variable, 'kind': c['kind'], 'expr': expr})
            added += 1
        put_rows(self.pom_tree, [(e['varname'], e['kind'], e['expr']) for e in self.pom_elements])
        self.pom_count.setText(f'{len(self.pom_elements)} elements in Page Objects')
        self.update_pom_code_preview()
        self.set_status(f'Added {added} element(s) to Page Objects; skipped {duplicates} duplicate(s).')

    def load_selected_pom_variable(self):
        rows = sorted({index.row() for index in self.pom_tree.selectedIndexes()})
        if len(rows) == 1 and rows[0] < len(self.pom_elements):
            self.pom_variable_entry.setText(self.pom_elements[rows[0]]['varname'])
            self.rename_pom_button.setEnabled(True)
        else:
            self.rename_pom_button.setEnabled(False)
            if not rows:
                self.pom_variable_entry.clear()

    def rename_selected_pom_variable(self):
        rows = sorted({index.row() for index in self.pom_tree.selectedIndexes()})
        if len(rows) != 1 or rows[0] >= len(self.pom_elements):
            self.set_status('Select one Page Object element to name.')
            return
        new_name = self.pom_variable_entry.text().strip()
        if not re.fullmatch(r'[A-Za-z_]\w*', new_name):
            self.set_status('Use a valid identifier: start with a letter or underscore, then use letters, digits, or underscores.')
            return
        index = rows[0]
        if any(i != index and item['varname'] == new_name for i, item in enumerate(self.pom_elements)):
            self.set_status('That variable name is already used in this Page Object.')
            return
        self.pom_elements[index]['varname'] = new_name
        put_rows(self.pom_tree, [(e['varname'], e['kind'], e['expr']) for e in self.pom_elements])
        self.pom_tree.selectRow(index)
        self.update_pom_code_preview()
        self.set_status(f'Variable renamed to {new_name}.')

    def update_pom_code_preview(self):
        if not hasattr(self, 'pom_code_text'):
            return
        class_name = self.pom_class_entry.text().strip() or 'LoginPage'
        self.pom_code_text.setPlainText(generate_pom_code(self.pom_elements, class_name, self.framework))

    def clear_pom(self):
        self.pom_elements.clear()
        put_rows(self.pom_tree, [])
        self.pom_count.setText('No elements added. Select locators from Element Picker.')
        self.update_pom_code_preview()
        self.set_status('Page Object element list cleared.')

    def copy_pom_class(self):
        if not self.pom_elements:
            self.set_status('Add elements to Page Objects first.')
            return
        self.copy(self.pom_code_text.toPlainText(), 'Page Object class copied.')

    def send_selected_to_auditor(self):
        candidates = self.selected_candidates()
        if not candidates:
            self.set_status('Select one or more locator rows to audit.')
            return
        existing = self.audit_input_text.toPlainText().splitlines()
        all_lines = {line.strip() for line in existing if line.strip()}
        added = 0
        for candidate in candidates:
            expr = candidate['expr']
            if expr not in all_lines:
                existing.append(expr)
                all_lines.add(expr)
                added += 1
        self.audit_input_text.setPlainText('\n'.join(existing).strip())
        self.switch_workspace(2)
        self.set_status(f'Sent {added} selector(s) to Health Audit.')

    def export_excel(self, candidates=None, context=None):
        if isinstance(candidates, bool):
            candidates = None
        export_rows = self.rows if candidates is None else candidates
        export_context = self.count_context if context is None else context
        if not export_rows:
            self.set_status('Generate locators before exporting.')
            return
        dest, selected_filter = QFileDialog.getSaveFileName(
            self, 'Export generated locators', 'Locators.xlsx',
            'Excel Workbook (*.xlsx);;CSV File (*.csv)')
        if not dest:
            return
        extension = Path(dest).suffix.lower()
        if extension not in ('.xlsx', '.csv'):
            dest += '.csv' if 'CSV' in selected_filter else '.xlsx'
            extension = Path(dest).suffix.lower()
        if extension == '.xlsx' and not HAS_OPENPYXL:
            self.popup('Dependency missing', 'Excel export requires openpyxl. Install: pip install openpyxl')
            return
        headers = ['#', 'Type', 'Strategy', 'Locator Expression', 'Stability Rank',
                   'Score Value', 'Matches', f'Code Binding ({self.framework})', 'Notes', 'Match Source']
        data = [[i, c['kind'], c['strategy'], c['expr'], c['rank'], c['score'],
                 c['count'] if c['count'] is not None else 'N/A',
                 format_framework_code(c, self.framework), c['note'], export_context]
                for i, c in enumerate(export_rows, 1)]
        try:
            if extension == '.xlsx':
                wb = openpyxl.Workbook()
                ws = wb.active
                ws.title = 'Locators'
                ws.sheet_view.showGridLines = False
                ws.append(headers)
                for row in data: ws.append(row)
                # Explicit cell colors keep the export readable regardless of Excel's
                # dark-mode or system theme settings.
                header_fill = PatternFill('solid', fgColor='FFFFFFFF')
                odd_fill = PatternFill('solid', fgColor='FFFFFFFF')
                even_fill = PatternFill('solid', fgColor='FFFFFFFF')
                thin_gray = Side(style='thin', color='FFD9E2F3')
                border = Border(bottom=thin_gray)
                for cell in ws[1]:
                    cell.fill = header_fill
                    cell.font = Font(name='Calibri', size=11, bold=True, color='FF1F2937')
                    cell.alignment = Alignment(vertical='center', horizontal='left', wrap_text=True)
                ws.row_dimensions[1].height = 30

                widths = [7, 12, 27, 52, 17, 13, 12, 62, 48, 24]
                for column, width in enumerate(widths, 1):
                    ws.column_dimensions[get_column_letter(column)].width = width
                wrap_columns = {3, 4, 8, 9, 10}
                for row_index in range(2, ws.max_row + 1):
                    fill = odd_fill if row_index % 2 == 0 else even_fill
                    max_lines = 1
                    for column_index in range(1, ws.max_column + 1):
                        cell = ws.cell(row_index, column_index)
                        cell.fill = fill
                        cell.font = Font(name='Calibri', size=10, color='FF1F2937')
                        cell.border = border
                        cell.alignment = Alignment(
                            vertical='top',
                            horizontal='center' if column_index in (1, 6, 7) else 'left',
                            wrap_text=column_index in wrap_columns,
                        )
                        if column_index in wrap_columns:
                            max_lines = max(max_lines, min(8, (len(str(cell.value or '')) // widths[column_index - 1]) + 1))
                    ws.row_dimensions[row_index].height = min(96, max(24, 15 * max_lines))

                ws.freeze_panes = 'A2'
                ws.auto_filter.ref = ws.dimensions
                ws.sheet_properties.pageSetUpPr.fitToPage = True
                ws.page_setup.orientation = 'landscape'
                ws.page_setup.fitToWidth = 1
                ws.page_setup.fitToHeight = 0
                wb.save(dest)
            else:
                with open(dest, 'w', newline='', encoding='utf-8-sig') as handle:
                    writer = csv.writer(handle)
                    writer.writerow(headers)
                    writer.writerows(data)
            self.set_status(f'Exported {len(export_rows)} locators to {dest}.')
        except Exception as exc:
            self.popup('Export failed', str(exc))

    def set_busy(self, active, message=None):
        self.busy = active
        self.progress.setVisible(active)
        for name in ('generate_button', 'scan_button', 'audit_button', 'test_button',
                     'picker_scan_button', 'picker_test_button', 'picker_generate_button'):
            button = getattr(self, name, None)
            if button is not None:
                button.setEnabled(not active)
        for name in ('scan_button', 'audit_button', 'test_button', 'picker_test_button'):
            button = getattr(self, name, None)
            if button is not None:
                button.setEnabled(not active and HAS_SELENIUM)
        picker_scan = getattr(self, 'picker_scan_button', None)
        if picker_scan is not None:
            picker_scan.setEnabled(not active and (HAS_SELENIUM or self._picker_loaded))
        if message:
            self.set_status(message)

    def launch_job(self, task, success, label_text):
        if self.busy:
            return
        self._job_success = success
        self.set_busy(True, label_text)
        thread = QThread(self)
        worker = Worker(task)
        worker.moveToThread(thread)
        record = (thread, worker)
        self.worker_threads.append(record)
        thread.started.connect(worker.run)
        worker.result.connect(self.receive_job_result)
        worker.error.connect(self.job_error)
        worker.finished.connect(self.job_finished)
        worker.finished.connect(thread.quit)
        worker.finished.connect(worker.deleteLater)
        thread.finished.connect(thread.deleteLater)
        thread.finished.connect(lambda: self.worker_threads.remove(record) if record in self.worker_threads else None)
        thread.start()

    def receive_job_result(self, payload):
        callback = self._job_success
        self._job_success = None
        if callback:
            callback(payload)

    def scan_whole_page(self):
        if not HAS_SELENIUM:
            self.set_status('Install Selenium to scan pages.')
            return
        try:
            url = validate_http_url(self.scan_url.text())
        except ValueError as exc:
            self.set_status(str(exc))
            return
        if not url:
            self.set_status('Enter a URL to scan.')
            return
        self.live_test_url.setText(url)
        excluded_types, excluded_strategies = self.get_excluded_locator_options()
        self.launch_job(lambda: scan_page_job(url, excluded_types, excluded_strategies), self.populate_scan_results,
                        'Scanning page source with Selenium…')

    def populate_scan_results(self, payload):
        self.loaded_bundle_id = None
        self.rows = payload['rows']
        self.count_context = 'Page-source snapshot'
        self.scan_summary.setText(
            f'Last scanned page:\n{payload["url"]}\n\n'
            f'Generated {len(self.rows)} top locator(s). Their match counts come from the captured page source; '
            'use Test live to verify them against the loaded page.'
        )
        self.show_candidates(f'Scanned {payload["url"]} · {len(self.rows)} top locators')

    def scan_picker_whole_page(self):
        try:
            url = validate_http_url(self.picker_url_entry.text())
        except ValueError as exc:
            self.picker_status_label.setText(str(exc))
            return
        self.picker_url_entry.setText(url)
        excluded_types, excluded_strategies = self.get_excluded_locator_options()

        if self._picker_loaded and self.picker_view:
            current_url = self.picker_view.url().toString()
            if current_url and current_url != 'about:blank':
                self.picker_status_label.setText('Reading the open page and generating locator candidates…')

                def process_markup(markup):
                    if not markup or not markup.strip():
                        self.picker_status_label.setText(
                            'The open page returned no HTML. Reload it or scan after the page finishes loading.'
                        )
                        return
                    self.launch_job(
                        lambda: scan_html_snapshot(markup, current_url, excluded_types, excluded_strategies),
                        self.populate_picker_scan_results,
                        'Scanning the currently open Element Picker page…',
                    )

                self.picker_view.page().toHtml(process_markup)
                return

        if not HAS_SELENIUM:
            self.picker_status_label.setText(
                'Open the page with Go to scan it in the embedded browser, or install Selenium for background scans.'
            )
            return
        self.launch_job(
            lambda: scan_page_job(url, excluded_types, excluded_strategies),
            self.populate_picker_scan_results,
            'Scanning the Element Picker URL…',
        )

    def generate_picker_fallback(self):
        if BeautifulSoup is None:
            self.popup('Dependency missing', 'Install BeautifulSoup: pip install beautifulsoup4')
            return
        markup = self.picker_generator_html.toPlainText().strip()
        if not markup:
            self.picker_generator_html.setFocus()
            self.picker_status_label.setText('Paste HTML before generating locators.')
            return
        excluded_types, excluded_strategies = self.get_excluded_locator_options()
        self.launch_job(
            lambda: generate_locators_job(markup, excluded_types, excluded_strategies),
            lambda payload, source=markup: self.populate_picker_input_results(payload, source),
            'Generating picker locators from pasted HTML…',
        )

    def populate_picker_input_results(self, payload, markup):
        capture_id = self._next_capture_id
        self._next_capture_id += 1
        capture_number = self._next_capture_number
        self._next_capture_number += 1
        bundle = {
            'id': capture_id,
            'number': capture_number,
            'url': '',
            'html': markup,
            'rows': payload['rows'],
            'status': f'{len(payload["rows"])} locator(s)',
            'source': 'picker_html',
        }
        bundle.update(self._element_name(markup))
        self._append_ready_picker_bundle(bundle, 'HTML input', markup)
        self.picker_detail_tabs.setCurrentIndex(1)
        self.picker_status_label.setText(
            f'Generated {len(bundle["rows"])} locator(s) from pasted HTML. Results are in Candidate Locators.'
        )
        self.set_status(f'Generated picker locators for <{payload["tag"]}> · {len(bundle["rows"])} candidates')

    def _append_ready_picker_bundle(self, bundle, source_label, tooltip=''):
        self.captured_bundles.append(bundle)
        capture_id = bundle['id']
        item = QTreeWidgetItem([
            f'#{bundle["number"]}  {bundle["display"]}', bundle['status'], source_label
        ])
        item.setFlags(item.flags() | Qt.ItemFlag.ItemIsUserCheckable)
        item.setCheckState(0, Qt.CheckState.Checked)
        item.setData(0, Qt.ItemDataRole.UserRole, capture_id)
        if tooltip:
            item.setToolTip(0, tooltip)
        for locator_index, locator in enumerate(bundle['rows']):
            child = QTreeWidgetItem([locator['expr'], locator['kind'], locator['strategy']])
            child.setData(0, Qt.ItemDataRole.UserRole, capture_id)
            child.setData(0, Qt.ItemDataRole.UserRole + 1, locator_index)
            child.setToolTip(0, locator['expr'])
            child.setToolTip(1, f'Score: {locator["score"]}/100 · {locator["rank"]}')
            child.setToolTip(2, locator['note'])
            item.addChild(child)
        self.captured_bundle_list.addTopLevelItem(item)
        self.on_capture_bundle_check_changed(item, 0)
        self.captured_bundle_list.setCurrentItem(item)
        self.update_picker_candidate_view(item)
        self.update_open_bundle_button()
        return item

    def populate_picker_scan_results(self, payload):
        if not payload.get('rows'):
            self.picker_status_label.setText(
                f'Page scan found {payload.get("elements", 0)} eligible element(s), but no locators remained. '
                'Check Ignore patterns or try a different page.'
            )
            self.set_status('Page scan finished with no locator candidates.')
            return
        capture_id = self._next_capture_id
        self._next_capture_id += 1
        capture_number = self._next_capture_number
        self._next_capture_number += 1
        host = urlparse(payload['url']).netloc or payload['url']
        bundle = {
            'id': capture_id,
            'number': capture_number,
            'url': payload['url'],
            'html': '',
            'rows': payload['rows'],
            'status': f'{len(payload["rows"])} locator(s)',
            'display': f'Page scan · {host}',
            'variable_prefix': 'page_scan',
            'source': 'page_scan',
        }
        self._append_ready_picker_bundle(bundle, 'Page scan', payload['url'])
        self.picker_detail_tabs.setCurrentIndex(0)
        self.picker_status_label.setText(
            f'Page scan complete: {len(bundle["rows"])} locator(s) from '
            f'{payload.get("elements", 0)} eligible element(s), added as a captured bundle.'
        )
        self.set_status(f'Scanned {payload["url"]} · {len(bundle["rows"])} locator(s)')

    def test_live_url(self):
        if not HAS_SELENIUM:
            self.set_status('Install Selenium to test selectors.')
            return
        candidates = self.selected_candidates()
        if not candidates:
            self.popup('Select locators', 'Check one or more locator rows before running a live test.')
            return
        try:
            url = validate_http_url(self.live_test_url.text())
        except ValueError as exc:
            self.set_status(str(exc))
            return
        if not url:
            self.set_status('Enter a target URL in the browser picker or Scan Page.')
            return
        selected = [candidate.copy() for candidate in candidates]
        self.launch_job(
            lambda: test_live_job(url, selected),
            self.show_live_test_results,
            f'Testing {len(selected)} selected locator(s) against live URL…',
        )

    def show_live_test_results(self, payload):
        results = payload['results']
        url = payload['url']
        found = sum(matches is not None and matches > 0 for _, matches, error in results if error is None)
        no_matches = sum(matches == 0 for _, matches, error in results if error is None)
        errors = sum(error is not None for _, _, error in results)
        summary = (f'Live test complete for {len(results)} locator(s).\n'
                   f'Found: {found}   |   No matches: {no_matches}   |   Errors: {errors}')
        lines = []
        for expr, matches, error in results:
            state = f'{matches} match(es)' if error is None else f'Evaluation error: {error}'
            lines.append(f'{expr}\n  {state}')
        dialog = QMessageBox(self)
        dialog.setIcon(QMessageBox.Icon.Information if errors == 0 else QMessageBox.Icon.Warning)
        dialog.setWindowTitle('Live locator test results')
        dialog.setText(summary)
        dialog.setInformativeText(url)
        dialog.setDetailedText('\n\n'.join(lines))
        dialog.setStandardButtons(QMessageBox.StandardButton.Ok)
        dialog.exec()
        self.set_status(f'Live locator test finished: {found} found, {no_matches} with no matches, {errors} errors.')

    def on_job_error(self, message):
        if self._active_capture_id is not None:
            bundle = next((item for item in self.captured_bundles if item['id'] == self._active_capture_id), None)
            if bundle:
                bundle['status'] = 'Generation failed'
                item = self._find_capture_item(bundle['id'])
                if item:
                    item.setText(1, 'Generation failed')
        self.popup('Operation failed', message)
        self.set_status('Operation failed: ' + message[:120])

    def on_job_finished(self):
        self.set_busy(False)
        self._active_capture_id = None
        self.process_next_capture()

    def run_health_audit_and_heal(self):
        if not HAS_SELENIUM:
            self.set_status('Install Selenium to run live Health Audit.')
            return
        try:
            url = validate_http_url(self.audit_url_entry.text())
        except ValueError as exc:
            self.set_status(str(exc))
            return
        locators = extract_audit_locators(self.audit_input_text.toPlainText())
        if not url or not locators:
            self.set_status('Enter an audit URL and at least one selector or Selenium locator declaration.')
            return
        put_rows(self.audit_tree, [])
        put_rows(self.suggestion_tree, [])
        self.launch_job(lambda: audit_job(url, locators), self.populate_audit_results,
                        f'Running {len(locators)} live locator checks…')

    def populate_audit_results(self, payload):
        put_rows(self.audit_tree, payload['results'])
        put_rows(self.suggestion_tree, payload['suggestions'])
        found = sum(status == 'Found' for _, _, status in payload['results'])
        self.set_status(f'Health Audit finished: {found}/{len(payload["results"])} locators found on live URL; '
                        f'{len(payload["suggestions"])} unverified suggestions/notes.')

    def navigate_picker(self):
        if not self.picker_view:
            self.picker_status_label.setText('Embedded browser unavailable. Install PySide6 Qt WebEngine.')
            return
        try:
            url = validate_http_url(self.picker_url_entry.text())
        except ValueError as exc:
            self.picker_status_label.setText(str(exc))
            return
        self.picker_url_entry.setText(url)
        self._picker_page_open = True
        self.picker_scan_button.setEnabled(HAS_SELENIUM)
        self.picker_browser_panel.show()
        self.picker_pick_mode.setEnabled(True)
        self.picker_close_page_button.setEnabled(True)
        self.picker_back_button.setEnabled(True)
        self.picker_forward_button.setEnabled(True)
        self.picker_reload_button.setEnabled(True)
        extent = (self.picker_splitter.height()
                  if self.picker_splitter.orientation() == Qt.Orientation.Vertical
                  else self.picker_splitter.width())
        self.picker_page_stack.setCurrentWidget(self.picker_view)
        self.picker_display_title.setText('PAGE')
        details_visible = self.picker_detail_tabs.isVisible()
        browser_ratio = (0.58 if details_visible else 0.66) if self.picker_splitter.orientation() == Qt.Orientation.Horizontal else (0.50 if details_visible else 0.60)
        first = max(1, round(extent * browser_ratio))
        self.picker_splitter.setSizes([first, max(1, extent - first)])
        self.picker_view.setUrl(QUrl(url))
        self.picker_status_label.setText('Loading page…')

    def close_picker_page(self):
        if not self.picker_view:
            return
        self.picker_poll_timer.stop()
        if self._picker_loaded:
            self.picker_view.page().runJavaScript('window.__pickerEnabled = false;')
        self._picker_page_open = False
        self._picker_loaded = False
        self.picker_scan_button.setEnabled(HAS_SELENIUM)
        self._picker_poll_in_flight = False
        self.picker_pick_mode.setEnabled(False)
        self.picker_view.stop()
        self.picker_view.setUrl(QUrl('about:blank'))
        self.picker_page_stack.setCurrentWidget(self.picker_fallback_page)
        self.picker_display_title.setText('GENERATE LOCATORS')
        self.picker_close_page_button.setEnabled(False)
        self.picker_back_button.setEnabled(False)
        self.picker_forward_button.setEnabled(False)
        self.picker_reload_button.setEnabled(False)
        extent = (self.picker_splitter.height()
                  if self.picker_splitter.orientation() == Qt.Orientation.Vertical
                  else self.picker_splitter.width())
        ratio = (0.42 if self.picker_detail_tabs.isVisible() else 0.66) if self.picker_splitter.orientation() == Qt.Orientation.Horizontal else (0.48 if self.picker_detail_tabs.isVisible() else 0.60)
        first = max(1, round(extent * ratio))
        self.picker_splitter.setSizes([first, max(1, extent - first)])
        self.picker_status_label.setText('Page closed. The HTML locator generator is ready, or enter a URL and click Go again.')

    def on_picker_url_changed(self, url):
        value = url.toString()
        if value and value != 'about:blank':
            self.picker_url_entry.setText(value)
            # Keep legacy hidden-workspace fields in sync only if that page exists.
            for field_name in ('live_test_url', 'scan_url'):
                field = getattr(self, field_name, None)
                if field is not None:
                    field.setText(value)

    def on_picker_load_finished(self, success):
        if (not self._picker_page_open or not self.picker_view
                or self.picker_view.url().toString() == 'about:blank'):
            return
        if not success:
            self._picker_loaded = False
            self.picker_status_label.setText('Page failed to load. Check the URL or connection and try again.')
            self.picker_poll_timer.stop()
            self._picker_page_open = False
            self.picker_scan_button.setEnabled(HAS_SELENIUM)
            self.picker_pick_mode.setEnabled(False)
            self.picker_page_stack.setCurrentWidget(self.picker_fallback_page)
            self.picker_display_title.setText('GENERATE LOCATORS')
            self.picker_close_page_button.setEnabled(False)
            self.picker_back_button.setEnabled(False)
            self.picker_forward_button.setEnabled(False)
            self.picker_reload_button.setEnabled(False)
            extent = (self.picker_splitter.height()
                      if self.picker_splitter.orientation() == Qt.Orientation.Vertical
                      else self.picker_splitter.width())
            ratio = (0.42 if self.picker_detail_tabs.isVisible() else 0.66) if self.picker_splitter.orientation() == Qt.Orientation.Horizontal else (0.48 if self.picker_detail_tabs.isVisible() else 0.60)
            first = max(1, round(extent * ratio))
            self.picker_splitter.setSizes([first, max(1, extent - first)])
            return
        self._picker_loaded = True
        self.picker_scan_button.setEnabled(True)
        self.picker_status_label.setText('Page ready. Click an element to capture it. Each click creates a separate bundle.')
        self.install_picker_script()
        if self.picker_pick_mode.isChecked():
            self.picker_poll_timer.start()

    def install_picker_script(self):
        if not self.picker_view or not self._picker_loaded:
            return
        enabled = 'true' if self.picker_pick_mode.isChecked() else 'false'
        self.picker_view.page().runJavaScript(PICKER_SCRIPT.replace('PICKER_ENABLED', enabled))

    def update_pick_mode(self, checked):
        if not self.picker_view or not HAS_WEBENGINE:
            return
        if checked and self._picker_loaded:
            self.install_picker_script()
            self.picker_poll_timer.start()
        else:
            self.picker_poll_timer.stop()
            if self._picker_loaded:
                self.picker_view.page().runJavaScript('window.__pickerEnabled = false;')
        self.picker_status_label.setText('Click to capture is on.' if checked else 'Browsing mode. Turn on Click to capture to select elements.')

    def poll_picker_selection(self):
        if not self.picker_view or not self._picker_loaded or self._picker_poll_in_flight or not self.picker_pick_mode.isChecked():
            return
        self._picker_poll_in_flight = True
        # Qt WebEngine evaluates JavaScript at page scope (unlike Selenium's
        # execute_script wrapper), so return must be inside an expression scope.
        js = '(function(){var x = window.__selectedElementHtml; window.__selectedElementHtml = null; return x;})()'
        self.picker_view.page().runJavaScript(js, self.on_picker_selection_polled)

    def on_picker_selection_polled(self, html_content):
        self._picker_poll_in_flight = False
        if not html_content:
            return
        capture_id = self._next_capture_id
        self._next_capture_id += 1
        capture_number = self._next_capture_number
        self._next_capture_number += 1
        bundle = {
            'id': capture_id,
            'number': capture_number,
            'url': self.picker_view.url().toString() if self.picker_view else '',
            'html': str(html_content),
            'rows': [],
            'status': 'Queued…',
        }
        bundle.update(self._element_name(bundle['html']))
        self.captured_bundles.append(bundle)
        self._pending_capture_ids.append(capture_id)
        item = QTreeWidgetItem([f'#{capture_number}  {bundle["display"]}', 'Generating…', ''])
        item.setFlags(item.flags() | Qt.ItemFlag.ItemIsUserCheckable)
        item.setCheckState(0, Qt.CheckState.Checked)
        item.setData(0, Qt.ItemDataRole.UserRole, capture_id)
        item.setToolTip(0, bundle['url'] + '\n\n' + bundle['html'])
        self.captured_bundle_list.addTopLevelItem(item)
        self.on_capture_bundle_check_changed(item, 0)
        self.captured_bundle_list.setCurrentItem(item)
        self.picker_status_label.setText(f'Captured element #{capture_number}; generating its independent locator bundle…')
        self.process_next_capture()

    def process_next_capture(self):
        if self.busy or not self._pending_capture_ids:
            return
        capture_id = self._pending_capture_ids.pop(0)
        bundle = next((item for item in self.captured_bundles if item['id'] == capture_id), None)
        if bundle is None:
            self.process_next_capture()
            return
        bundle['status'] = 'Generating…'
        item = self._find_capture_item(capture_id)
        if item:
            item.setText(1, 'Generating…')
        self._active_capture_id = capture_id
        excluded_types, excluded_strategies = self.get_excluded_locator_options()
        self.launch_job(
            lambda markup=bundle['html'], types=tuple(excluded_types), strategies=tuple(excluded_strategies):
                generate_locators_job(markup, types, strategies),
            lambda payload, cid=capture_id: self.finish_capture_generation(cid, payload),
            f'Generating locators for captured element #{bundle["number"]}…',
        )

    @staticmethod
    def _element_name(markup):
        soup = BeautifulSoup(markup, 'html.parser') if BeautifulSoup else None
        element = soup.find(True) if soup else None
        if not element:
            return {'display': 'Captured element', 'variable_prefix': 'captured_element'}
        identity = next((element.get(key) for key in (
            'data-testid', 'data-test', 'data-qa', 'aria-label', 'name', 'id',
            'placeholder', 'title') if element.get(key)), '')
        if not identity:
            identity = ' '.join(element.stripped_strings)
        identity = re.sub(r'\s+', ' ', str(identity)).strip()[:48]
        display = f'<{element.name}>' + (f' · {identity}' if identity else ' · clicked element')
        variable = re.sub(r'[^A-Za-z0-9]+', '_', f'{element.name}_{identity}').strip('_').lower()
        if not variable:
            variable = 'captured_element'
        if variable[0].isdigit():
            variable = 'element_' + variable
        return {'display': display, 'variable_prefix': variable}

    @classmethod
    def _element_description(cls, markup):
        return cls._element_name(markup)['display']

    def _find_capture_item(self, capture_id):
        for index in range(self.captured_bundle_list.topLevelItemCount()):
            item = self.captured_bundle_list.topLevelItem(index)
            if item.data(0, Qt.ItemDataRole.UserRole) == capture_id:
                return item
        return None

    def finish_capture_generation(self, capture_id, payload):
        bundle = next((item for item in self.captured_bundles if item['id'] == capture_id), None)
        if bundle is None:
            return
        bundle['rows'] = payload['rows']
        bundle['status'] = f'{len(payload["rows"])} locator(s)'
        item = self._find_capture_item(capture_id)
        if item:
            item.setText(0, f'#{bundle["number"]}  {bundle["display"]}')
            item.setText(1, bundle['status'])
            item.takeChildren()
            for locator_index, locator in enumerate(bundle['rows']):
                child = QTreeWidgetItem([
                    locator['expr'], locator['kind'], locator['strategy']
                ])
                child.setData(0, Qt.ItemDataRole.UserRole, capture_id)
                child.setData(0, Qt.ItemDataRole.UserRole + 1, locator_index)
                child.setToolTip(0, locator['expr'])
                child.setToolTip(1, f'Score: {locator["score"]}/100 · {locator["rank"]}')
                child.setToolTip(2, locator['note'])
                item.addChild(child)
            item.setExpanded(False)
        self.update_picker_candidate_view(item)
        self.load_capture_bundle(bundle, switch_workspace=False)
        self.update_open_bundle_button()
        self.set_status(f'Captured {bundle["display"]}: generated {len(bundle["rows"])} locator(s) in bundle #{bundle["number"]}.')
        if self.captured_bundle_list.currentItem() is item:
            self.picker_status_label.setText(f'{bundle["display"]} ready: {len(bundle["rows"])} locator(s). Its candidates are available in Element Picker; the next click creates a separate bundle.')

    def update_picker_candidate_view(self, current=None, _previous=None):
        bundles = self.checked_capture_bundles()
        rows = [(bundle, locator_index, locator)
                for bundle in bundles
                for locator_index, locator in enumerate(bundle['rows'])]
        self.picker_candidate_sources = []
        self.picker_candidate_tree.blockSignals(True)
        try:
            self.picker_candidate_tree.setRowCount(0)
            for bundle, locator_index, candidate in rows:
                row = self.picker_candidate_tree.rowCount()
                self.picker_candidate_tree.insertRow(row)
                self.picker_candidate_sources.append((bundle['id'], locator_index))
                check = QTableWidgetItem('')
                check.setFlags(check.flags() | Qt.ItemFlag.ItemIsUserCheckable | Qt.ItemFlag.ItemIsEnabled)
                single_mode = self.picker_selection_mode.currentText() == 'Single selection'
                check.setCheckState(
                    Qt.CheckState.Checked if (not single_mode or row == 0)
                    else Qt.CheckState.Unchecked
                )
                self.picker_candidate_tree.setItem(row, 0, check)
                values = (candidate['expr'], candidate['kind'], candidate['strategy'],
                          f'{candidate["rank"]} · {candidate["score"]}/100',
                          str(candidate['count']) if candidate.get('count') is not None else 'N/A',
                          f'#{bundle["number"]} {bundle["display"]}')
                for column, value in enumerate(values, 1):
                    cell = QTableWidgetItem(str(value))
                    if column in (1, 6):
                        cell.setToolTip(str(value))
                    self.picker_candidate_tree.setItem(row, column, cell)
        finally:
            self.picker_candidate_tree.blockSignals(False)
        if bundles:
            self.picker_candidates_caption.setText(
                f'{len(bundles)} captured bundle(s) selected · {len(rows)} candidate locator(s)'
            )
        else:
            self.picker_candidates_caption.setText(
                'Select one or more captured bundles to show their candidates and code preview.'
            )
        self.picker_select_all.blockSignals(True)
        self.picker_select_all.setChecked(
            bool(rows) and self.picker_selection_mode.currentText() == 'Multiple selection'
        )
        self.picker_select_all.blockSignals(False)
        self.update_picker_code_preview()
        selected_item = current or self.captured_bundle_list.currentItem()
        selected_parent = selected_item.parent() if selected_item and selected_item.parent() else selected_item
        preferred_source = None
        if selected_parent:
            selected_id = selected_parent.data(0, Qt.ItemDataRole.UserRole)
            if selected_item.parent():
                preferred_source = (selected_id, selected_item.data(0, Qt.ItemDataRole.UserRole + 1))
            else:
                preferred_source = next((source for source in self.picker_candidate_sources
                                         if source[0] == selected_id), None)
        row_to_select = next((index for index, source in enumerate(self.picker_candidate_sources)
                              if source == preferred_source), None)
        if row_to_select is None and self.picker_candidate_sources:
            row_to_select = 0
        if row_to_select is not None:
            self.picker_candidate_tree.selectRow(row_to_select)
        else:
            self.picker_locator_details.setPlainText(
                'Select one or more captured bundles to inspect locator details.'
            )

    def update_picker_locator_details(self):
        entries = self.selected_picker_candidate_entries()
        if not entries:
            self.picker_locator_details.setPlainText(
                'Check one or more candidate locators to inspect their details.'
            )
            return
        sections = []
        for number, (candidate, bundle, locator_index) in enumerate(entries, 1):
            matches = candidate['count'] if candidate.get('count') is not None else 'N/A'
            live = self._picker_live_results.get((bundle['id'], locator_index))
            if live is None:
                live_text = 'Live test: Not tested yet.'
            elif live.get('error'):
                live_text = f'Live test: Evaluation error — {live["error"]}'
            elif live.get('matches', 0) > 0:
                live_text = f'Live test: Found {live["matches"]} match(es) at {live["url"]}.'
            else:
                live_text = f'Live test: No matches at {live["url"]}.'
            sections.append(
                f'[{number}] Captured element: #{bundle["number"]} {bundle["display"]}\n'
                f'Locator: {candidate["expr"]}\n'
                f'Type: {candidate["kind"]}   Strategy: {candidate["strategy"]}\n'
                f'Score: {candidate["score"]}/100 ({candidate["rank"]})\n'
                f'Snapshot matches: {matches}\n'
                f'{live_text}\n\n'
                f'{candidate["note"]}\n\n'
                f'Framework binding ({self.framework}):\n'
                f'{format_framework_code(candidate, self.framework)}'
            )
        self.picker_locator_details.setPlainText('\n\n' + ('\n\n' + '—' * 48 + '\n\n').join(sections))

    def selected_picker_candidates(self):
        return [candidate for candidate, _bundle, _locator_index in self.selected_picker_candidate_entries()]

    def selected_picker_candidate_entries(self):
        entries = []
        for row in range(self.picker_candidate_tree.rowCount()):
            check = self.picker_candidate_tree.item(row, 0)
            if check and check.checkState() == Qt.CheckState.Checked:
                if row >= len(self.picker_candidate_sources):
                    continue
                capture_id, locator_index = self.picker_candidate_sources[row]
                bundle = next((entry for entry in self.captured_bundles
                               if entry['id'] == capture_id), None)
                if bundle and 0 <= locator_index < len(bundle['rows']):
                    entries.append((bundle['rows'][locator_index], bundle, locator_index))
        return entries

    def add_picker_candidates_to_pom(self):
        entries = self.selected_picker_candidate_entries()
        if not entries:
            self.set_status('Select one or more picker locator rows to add to Page Objects.')
            return
        added = duplicates = 0
        for candidate, bundle, _locator_index in entries:
            expr = candidate['expr']
            if any(element['expr'] == expr for element in self.pom_elements):
                duplicates += 1
                continue
            name = generate_descriptive_name(expr)
            prefix = bundle.get('variable_prefix', '')
            base = f'{prefix}_{name}' if prefix else name
            variable = base
            suffix = 1
            existing = {element['varname'] for element in self.pom_elements}
            while variable in existing:
                variable = f'{base}{suffix}'
                suffix += 1
            self.pom_elements.append({'varname': variable, 'kind': candidate['kind'], 'expr': expr})
            added += 1
        put_rows(self.pom_tree, [(element['varname'], element['kind'], element['expr'])
                                 for element in self.pom_elements])
        self.pom_count.setText(f'{len(self.pom_elements)} elements in Page Objects')
        self.update_pom_code_preview()
        self.set_status(f'Added {added} picker locator(s) to Page Objects; skipped {duplicates} duplicate(s).')

    def send_picker_candidates_to_audit(self):
        candidates = self.selected_picker_candidates()
        if not candidates:
            self.set_status('Select one or more picker locator rows to send to Health Audit.')
            return
        existing = self.audit_input_text.toPlainText().splitlines()
        seen = {line.strip() for line in existing if line.strip()}
        added = 0
        for candidate in candidates:
            expression = candidate['expr']
            if expression not in seen:
                existing.append(expression)
                seen.add(expression)
                added += 1
        self.audit_input_text.setPlainText('\n'.join(existing).strip())
        self.switch_workspace(2)
        self.set_status(f'Sent {added} picker locator(s) to Health Audit.')

    def test_picker_candidates_live(self):
        if not HAS_SELENIUM:
            self.set_status('Install Selenium to test selectors.')
            return
        entries = self.selected_picker_candidate_entries()
        if not entries:
            self.popup('Select locators', 'Check one or more locator rows before running a live test.')
            return
        try:
            url = validate_http_url(self.picker_url_entry.text())
        except ValueError as exc:
            self.set_status(str(exc))
            return
        if not url:
            self.set_status('Enter a URL in the Element Picker address bar.')
            return
        selected = [candidate.copy() for candidate, _bundle, _index in entries]
        self.launch_job(
            lambda: test_live_job(url, selected),
            lambda payload, entries=entries: self.show_picker_live_test_results(payload, entries),
            f'Testing {len(selected)} picker locator(s) against live URL…',
        )

    def show_picker_live_test_results(self, payload, entries):
        for (candidate, bundle, locator_index), result in zip(entries, payload.get('results', [])):
            _expression, matches, error = result
            self._picker_live_results[(bundle['id'], locator_index)] = {
                'matches': matches, 'error': error, 'url': payload.get('url', '')
            }
        self.update_picker_locator_details()
        self.show_live_test_results(payload)

    def export_picker_candidates(self):
        candidates = self.selected_picker_candidates()
        if not candidates:
            self.set_status('Select one or more picker locator rows to export.')
            return
        selected_bundles = self.checked_capture_bundles()
        source = ', '.join(f'#{bundle["number"]} {bundle["display"]}' for bundle in selected_bundles)
        self.export_excel(candidates, f'Element Picker candidates · {source or "page scan"}')

    def on_picker_candidate_changed(self, item):
        if item.column() != 0:
            return
        if (self.picker_selection_mode.currentText() == 'Single selection'
                and item.checkState() == Qt.CheckState.Checked):
            self.picker_candidate_tree.blockSignals(True)
            try:
                for row in range(self.picker_candidate_tree.rowCount()):
                    other = self.picker_candidate_tree.item(row, 0)
                    if other is not None and other is not item:
                        other.setCheckState(Qt.CheckState.Unchecked)
            finally:
                self.picker_candidate_tree.blockSignals(False)
        all_checked = (self.picker_selection_mode.currentText() == 'Multiple selection'
                       and bool(self.picker_candidate_tree.rowCount()) and all(
            self.picker_candidate_tree.item(row, 0).checkState() == Qt.CheckState.Checked
            for row in range(self.picker_candidate_tree.rowCount())
        ))
        self.picker_select_all.blockSignals(True)
        self.picker_select_all.setChecked(all_checked)
        self.picker_select_all.blockSignals(False)
        self.update_picker_code_preview()
        self.update_picker_locator_details()

    def set_picker_selection_mode(self, mode):
        single = mode == 'Single selection'
        self.picker_select_all.setEnabled(not single)
        if single:
            self.picker_candidate_tree.blockSignals(True)
            try:
                kept_one = False
                for row in range(self.picker_candidate_tree.rowCount()):
                    item = self.picker_candidate_tree.item(row, 0)
                    if item is None or item.checkState() != Qt.CheckState.Checked:
                        continue
                    if kept_one:
                        item.setCheckState(Qt.CheckState.Unchecked)
                    else:
                        kept_one = True
            finally:
                self.picker_candidate_tree.blockSignals(False)
        self.picker_select_all.blockSignals(True)
        self.picker_select_all.setChecked(
            not single and bool(self.picker_candidate_tree.rowCount()) and all(
                self.picker_candidate_tree.item(row, 0).checkState() == Qt.CheckState.Checked
                for row in range(self.picker_candidate_tree.rowCount())
            )
        )
        self.picker_select_all.blockSignals(False)
        self.update_picker_code_preview()
        self.update_picker_locator_details()

    def set_all_picker_candidates(self, checked):
        single = self.picker_selection_mode.currentText() == 'Single selection'
        self.picker_candidate_tree.blockSignals(True)
        try:
            for row in range(self.picker_candidate_tree.rowCount()):
                item = self.picker_candidate_tree.item(row, 0)
                if item:
                    selected = checked and (not single or row == 0)
                    item.setCheckState(Qt.CheckState.Checked if selected else Qt.CheckState.Unchecked)
        finally:
            self.picker_candidate_tree.blockSignals(False)
        self.update_picker_code_preview()
        self.update_picker_locator_details()

    def update_picker_code_preview(self, *_):
        if not hasattr(self, 'picker_code_preview'):
            return
        candidates = self.selected_picker_candidates()
        code = '\n'.join(format_framework_code(candidate, self.framework) for candidate in candidates)
        self.picker_code_preview.setPlainText(code or 'Select one or more candidate locators to preview framework code.')

    def copy_picker_candidate_selectors(self):
        candidates = self.selected_picker_candidates()
        if candidates:
            self.copy('\n'.join(candidate['expr'] for candidate in candidates),
                      f'Copied {len(candidates)} picker locator(s).')
        else:
            self.set_status('Select one or more picker locator rows to copy.')

    def copy_picker_candidate_code(self):
        candidates = self.selected_picker_candidates()
        if candidates:
            self.copy('\n'.join(format_framework_code(candidate, self.framework) for candidate in candidates),
                      f'Copied {len(candidates)} {self.framework} picker binding(s).')
        else:
            self.set_status('Select one or more picker locator rows to copy code.')

    def _top_level_capture_items(self):
        return [self.captured_bundle_list.topLevelItem(index)
                for index in range(self.captured_bundle_list.topLevelItemCount())]

    def checked_capture_bundles(self):
        checked_ids = {
            item.data(0, Qt.ItemDataRole.UserRole)
            for item in self._top_level_capture_items()
            if item.checkState(0) == Qt.CheckState.Checked
        }
        return [bundle for bundle in self.captured_bundles if bundle['id'] in checked_ids]

    def on_capture_bundle_check_changed(self, item, column):
        if column != 0 or item.parent():
            return
        if (self.capture_selection_mode.currentText() == 'Single selection'
                and item.checkState(0) == Qt.CheckState.Checked):
            self.captured_bundle_list.blockSignals(True)
            try:
                for other in self._top_level_capture_items():
                    if other is not item:
                        other.setCheckState(0, Qt.CheckState.Unchecked)
            finally:
                self.captured_bundle_list.blockSignals(False)
        all_checked = (
            self.capture_selection_mode.currentText() == 'Multiple selection'
            and bool(self._top_level_capture_items())
            and all(item.checkState(0) == Qt.CheckState.Checked
                    for item in self._top_level_capture_items())
        )
        self.select_all_captures.blockSignals(True)
        self.select_all_captures.setChecked(all_checked)
        self.select_all_captures.blockSignals(False)
        self.update_open_bundle_button()
        self.update_picker_candidate_view()

    def set_capture_selection_mode(self, mode):
        single = mode == 'Single selection'
        self.select_all_captures.setEnabled(not single)
        if single:
            self.captured_bundle_list.blockSignals(True)
            try:
                items = self._top_level_capture_items()
                current = self.captured_bundle_list.currentItem()
                current = current.parent() if current and current.parent() else current
                keep = current if current and current.checkState(0) == Qt.CheckState.Checked else next(
                    (item for item in items if item.checkState(0) == Qt.CheckState.Checked), None
                )
                for item in items:
                    if item.checkState(0) != Qt.CheckState.Checked:
                        continue
                    if item is not keep:
                        item.setCheckState(0, Qt.CheckState.Unchecked)
            finally:
                self.captured_bundle_list.blockSignals(False)
        all_checked = (not single and bool(self._top_level_capture_items())
                       and all(item.checkState(0) == Qt.CheckState.Checked
                               for item in self._top_level_capture_items()))
        self.select_all_captures.blockSignals(True)
        self.select_all_captures.setChecked(all_checked)
        self.select_all_captures.blockSignals(False)
        self.update_picker_candidate_view()

    def set_all_captures_checked(self, checked):
        single = self.capture_selection_mode.currentText() == 'Single selection'
        self.captured_bundle_list.blockSignals(True)
        try:
            for index, item in enumerate(self._top_level_capture_items()):
                selected = checked and (not single or index == 0)
                item.setCheckState(0, Qt.CheckState.Checked if selected else Qt.CheckState.Unchecked)
        finally:
            self.captured_bundle_list.blockSignals(False)
        self.update_open_bundle_button()
        self.update_picker_candidate_view()

    def update_open_bundle_button(self, *_):
        selected = self.captured_bundle_list.currentItem() if hasattr(self, 'captured_bundle_list') else None
        parent = selected.parent() if selected and selected.parent() else selected
        capture_id = parent.data(0, Qt.ItemDataRole.UserRole) if parent else None
        bundle = next((item for item in self.captured_bundles if item['id'] == capture_id), None)
        ready = bool(bundle and bundle['status'] not in ('Generating…', 'Queued…', 'Generation failed'))
        self.open_capture_button.setEnabled(ready)
        selected_bundles = self.checked_capture_bundles()
        if not selected_bundles and bundle:
            selected_bundles = [bundle]
        can_copy = any(item['rows'] and item['status'] not in
                       ('Generating…', 'Queued…', 'Generation failed')
                       for item in selected_bundles)
        self.copy_capture_bundle_button.setEnabled(can_copy)
        self.copy_capture_code_button.setEnabled(can_copy)
        self.copy_capture_locator_button.setEnabled(
            ready and bool(selected and selected.parent())
        )

    def _selected_capture_bundle(self):
        selected = self.captured_bundle_list.currentItem()
        parent = selected.parent() if selected and selected.parent() else selected
        if not parent:
            return None, None
        capture_id = parent.data(0, Qt.ItemDataRole.UserRole)
        bundle = next((entry for entry in self.captured_bundles if entry['id'] == capture_id), None)
        return bundle, selected

    def on_capture_tree_double_clicked(self, item, _column):
        if item.parent():
            self.copy_selected_capture_locator()
        else:
            self.open_selected_capture()

    def on_capture_tree_clicked(self, item, _column):
        if item.parent():
            parent = item.parent()
            if parent.checkState(0) != Qt.CheckState.Checked:
                parent.setCheckState(0, Qt.CheckState.Checked)
            self.update_picker_candidate_view(item)
            return
        row_rect = self.captured_bundle_list.visualItemRect(item)
        click_x = self.captured_bundle_list.viewport().mapFromGlobal(QCursor.pos()).x()
        label_start = row_rect.left() + self.captured_bundle_list.indentation() + 42
        if click_x >= label_start and item.checkState(0) != Qt.CheckState.Checked:
            item.setCheckState(0, Qt.CheckState.Checked)
        # Clicking the row label behaves like a disclosure control; native
        # clicks on the tree arrow remain responsible for collapsing it.
        if click_x >= row_rect.left() + self.captured_bundle_list.indentation() + 28:
            item.setExpanded(not item.isExpanded())

    def copy_selected_capture_locator(self):
        bundle, selected = self._selected_capture_bundle()
        if not bundle or not selected or not selected.parent():
            return
        index = selected.data(0, Qt.ItemDataRole.UserRole + 1)
        if isinstance(index, int) and 0 <= index < len(bundle['rows']):
            locator = bundle['rows'][index]
            self.copy(locator['expr'], f'Copied locator from {bundle["display"]}.')

    def copy_capture_bundle_selectors(self):
        bundles = self.checked_capture_bundles()
        if not bundles:
            current, _selected = self._selected_capture_bundle()
            bundles = [current] if current else []
        locators = [locator for bundle in bundles for locator in bundle['rows']]
        if locators:
            self.copy('\n'.join(locator['expr'] for locator in locators),
                      f'Copied {len(locators)} selectors from {len(bundles)} captured element bundle(s).')

    def copy_capture_bundle_code(self):
        bundles = self.checked_capture_bundles()
        if not bundles:
            current, _selected = self._selected_capture_bundle()
            bundles = [current] if current else []
        locators = [locator for bundle in bundles for locator in bundle['rows']]
        if locators:
            code = '\n'.join(format_framework_code(locator, self.framework) for locator in locators)
            self.copy(code, f'Copied {len(locators)} {self.framework} bindings from {len(bundles)} captured element bundle(s).')

    def open_selected_capture(self, *_):
        bundle, _selected = self._selected_capture_bundle()
        if not bundle or bundle['status'] in ('Generating…', 'Queued…', 'Generation failed'):
            return
        self.load_capture_bundle(bundle, switch_workspace=True)

    def load_capture_bundle(self, bundle, switch_workspace):
        capture_id = bundle['id']
        self.loaded_bundle_id = capture_id
        if bundle.get('source') == 'page_scan':
            self.count_context = f'Page scan snapshot: {bundle["url"]}'
        else:
            self.count_context = f'Captured bundle #{bundle["number"]}: {bundle["display"]}'
        if switch_workspace:
            self.switch_workspace(0)
            self.picker_detail_tabs.setCurrentIndex(0)
        self.update_picker_candidate_view()
        self.set_status(f'Bundle #{bundle["number"]} · {bundle["display"]} · {len(bundle["rows"])} locators')

    def clear_captured_bundles(self):
        self.captured_bundles.clear()
        self._picker_live_results.clear()
        self._pending_capture_ids.clear()
        self.loaded_bundle_id = None
        self._active_capture_id = None
        self._next_capture_number = 1
        self.captured_bundle_list.clear()
        self.select_all_captures.blockSignals(True)
        self.select_all_captures.setChecked(False)
        self.select_all_captures.blockSignals(False)
        self.update_picker_candidate_view()
        self.update_open_bundle_button()
        self.picker_status_label.setText('Captured bundles cleared and numbering restarted. Click a page element to capture a new bundle.')
    def toggle_theme(self):
        self.dark = not self.dark
        self.apply_theme()

    def apply_theme(self):
        dark = self.dark
        scale = self.ui_scale
        px = lambda value: max(1, round(value * scale))
        bg = '#0f172a' if dark else '#f4f7fb'
        surface = '#151e30' if dark else '#ffffff'
        sidebar = '#111b2d' if dark else '#eaf0f8'
        muted = '#9aa9c0' if dark else '#5e6d83'
        fg = '#e9eef9' if dark else '#1d2c43'
        border = '#344258' if dark else '#d9e2ef'
        field = '#101a2b' if dark else '#f9fbfe'
        hover = '#233553' if dark else '#ddeafa'
        table_alt = '#1a263b' if dark else '#f6f9fe'
        self.theme_button.setText('Theme' if self.width() < 700 * scale else ('Light theme' if dark else 'Dark theme'))
        self.zoom_label.setText(f'{round(scale * 100)}%')
        self.progress.setFixedWidth(px(170))
        self.progress.setFixedHeight(px(9))
        self.sidebar.setFixedWidth(px(194))
        for button in self.nav_buttons:
            button.setMinimumHeight(px(40))
        if hasattr(self, 'generate_button'):
            self.generate_button.setMinimumHeight(px(39))
        self.setStyleSheet(f'''
            QWidget {{ background: {bg}; color: {fg}; font-family: "Segoe UI", "Arial"; font-size: {px(12)}px; }}
            QFrame#header {{ background: {surface}; border-bottom: 1px solid {border}; }}
            QFrame#sidebar {{ background: {sidebar}; border-right: 1px solid {border}; }}
            QFrame#footer {{ background: {surface}; border-top: 1px solid {border}; }}
            QFrame#panel {{ background: {surface}; border: 1px solid {border}; border-radius: 10px; }}
            QFrame#pickerCard {{ background: {sidebar}; border: 1px solid {border}; border-radius: 8px; }}
            QLabel {{ background: transparent; }}
            QLabel#appTitle {{ font-size: {px(18)}px; font-weight: 700; }}
            QLabel#pageTitle {{ font-size: {px(20)}px; font-weight: 700; }}
            QLabel#sectionTitle {{ font-size: {px(11)}px; font-weight: 700; letter-spacing: {px(1)}px; color: {muted}; }}
            QLabel#muted {{ color: {muted}; font-size: {px(11)}px; }}
            QPushButton {{ background: {surface}; border: 1px solid {border}; border-radius: 6px;
                          padding: {px(8)}px {px(11)}px; font-weight: 500; }}
            QPushButton:hover {{ background: {hover}; border-color: #3b82f6; }}
            QPushButton:focus {{ border: 2px solid #60a5fa; }}
            QPushButton:disabled {{ color: {muted}; border-color: {border}; }}
            QPushButton#primaryButton {{ background: #2563eb; color: white; border-color: #2563eb; font-weight: 700; }}
            QPushButton#primaryButton:hover {{ background: #1d4ed8; }}
            QPushButton#navButton {{ background: transparent; text-align: left; border: none; padding: {px(10)}px {px(13)}px; }}
            QPushButton#navButton:hover {{ background: {hover}; }}
            QPushButton#navButton:checked {{ background: #2563eb; color: white; }}
            QPushButton#modeButton {{ padding: {px(7)}px {px(5)}px; }}
            QPushButton#modeButton:checked {{ background: #2563eb; color: white; border-color: #2563eb; }}
            QPushButton#zoomButton {{ padding: {px(5)}px {px(8)}px; min-width: {px(24)}px; }}
            QLineEdit, QPlainTextEdit, QComboBox, QTableWidget {{
                background: {field}; color: {fg}; border: 1px solid {border}; border-radius: 6px;
                selection-background-color: #2563eb; selection-color: white; padding: {px(5)}px;
            }}
            QPlainTextEdit#codeArea {{ font-family: Consolas, "Cascadia Code", monospace; font-size: {px(12)}px; }}
            QTextBrowser#guideBrowser {{ background: {field}; color: {fg}; border: 1px solid {border};
                                        border-radius: 7px; padding: {px(10)}px; selection-background-color: #2563eb;
                                        selection-color: white; }}
            QPlainTextEdit:focus, QLineEdit:focus, QComboBox:focus, QTableWidget:focus {{ border: 2px solid #60a5fa; }}
            QTableWidget {{ alternate-background-color: {table_alt}; gridline-color: {border}; }}
            QTabWidget::pane {{ background: {surface}; border: 1px solid {border}; border-radius: 7px; top: -1px; }}
            QTabBar::tab {{ background: {sidebar}; color: {muted}; border: 1px solid {border};
                            padding: {px(8)}px {px(13)}px; margin-right: {px(3)}px; border-top-left-radius: 5px;
                            border-top-right-radius: 5px; }}
            QTabBar::tab:selected {{ background: {surface}; color: {fg}; border-bottom-color: {surface}; font-weight: 600; }}
            QTabBar::tab:hover:!selected {{ background: {hover}; color: {fg}; }}
            QHeaderView::section {{ background: {sidebar}; color: {fg}; border: none;
                                   border-bottom: 1px solid {border}; padding: {px(9)}px {px(6)}px; font-weight: 600; }}
            QTableWidget::item:selected {{ background: #2563eb; color: white; }}
            QComboBox::drop-down {{ border: none; width: 28px; background: transparent;
                                    subcontrol-origin: padding; subcontrol-position: right center; }}
            QComboBox::down-arrow {{ image: none; }}
            QComboBox QAbstractItemView {{ background: {surface}; color: {fg};
                border: 1px solid {border}; border-radius: 7px; padding: 5px;
                outline: none; selection-background-color: #2563eb; selection-color: white; }}
            QComboBox QAbstractItemView::item {{ min-height: {px(28)}px; padding: {px(5)}px {px(10)}px; border-radius: 4px; }}
            QComboBox QAbstractItemView::item:hover {{ background: {hover}; }}
            QToolButton#excludeButton {{ background: {field}; color: {fg}; border: 1px solid {border};
                                         border-radius: 6px; padding: {px(5)}px {px(30)}px {px(5)}px {px(10)}px; }}
            QToolButton#excludeButton::menu-indicator {{ image: none; width: 0px; }}
            QToolButton#excludeButton:hover {{ background: {hover}; border-color: #3b82f6; }}
            QToolButton#drawerButton {{ background: {surface}; color: {fg}; border: 1px solid {border};
                                        border-radius: 6px; padding: {px(6)}px; min-width: {px(30)}px; }}
            QToolButton#drawerButton:hover {{ background: {hover}; border-color: #3b82f6; }}
            QMenu {{ background: {surface}; color: {fg}; border: 1px solid {border}; padding: {px(5)}px; }}
            QMenu::item {{ padding: {px(6)}px {px(26)}px {px(6)}px {px(8)}px; border-radius: 4px; }}
            QMenu::item:selected {{ background: {hover}; }}
            QMenu::item:disabled {{ color: {muted}; padding-top: {px(8)}px; padding-bottom: {px(4)}px; }}
            QMenu::separator {{ height: 1px; background: {border}; margin: {px(4)}px {px(6)}px; }}
            QMenu::indicator {{ width: 14px; height: 14px; }}
            QCheckBox {{ spacing: {px(8)}px; }}
            QSplitter::handle {{ background: {border}; }}
            QSplitter::handle:horizontal {{ width: {px(5)}px; }}
            QSplitter::handle:vertical {{ height: {px(5)}px; }}
            QProgressBar {{ background: {border}; border: none; border-radius: 4px; }}
            QProgressBar::chunk {{ background: #2563eb; border-radius: 4px; }}
            QScrollBar:vertical, QScrollBar:horizontal {{ background: {field}; }}
            QScrollBar::handle:vertical, QScrollBar::handle:horizontal {{ background: {border}; border-radius: 4px; }}
        ''')
        if hasattr(self, 'help_browser'):
            self.help_browser.document().setDefaultStyleSheet(
                f'body {{ color: {fg}; font-family: "Segoe UI", Arial, sans-serif; }} '
                f'code {{ color: {fg}; background-color: {surface}; }} '
                f'hr {{ color: {border}; }}'
            )

    def closeEvent(self, event):
        if self.picker_poll_timer:
            self.picker_poll_timer.stop()
        if any(thread.isRunning() for thread, _ in self.worker_threads):
            self.set_status('A browser operation is running. Close after it finishes.')
            if not self._close_when_idle:
                self._close_when_idle = True
                for thread, worker in list(self.worker_threads):
                    thread.finished.connect(self._retry_close_when_idle)
            event.ignore()
            return
        event.accept()

    def _retry_close_when_idle(self):
        self._close_when_idle = False
        self.close()


def main():
    app = QApplication(sys.argv)
    app.setStyle('Fusion')
    window = MainWindow()
    window.show()
    return app.exec()


if __name__ == '__main__':
    sys.exit(main())

